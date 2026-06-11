from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook


def sheet_rows(ws) -> list[dict[str, str]]:
    values = list(ws.iter_rows(values_only=True))
    header = [str(cell) if cell is not None else "" for cell in values[0]]
    rows = []
    for raw_row in values[1:]:
        row = {}
        has_value = False
        for key, value in zip(header, raw_row):
            normalized = "" if value is None else str(value)
            row[key] = normalized
            has_value = has_value or bool(normalized)
        if has_value:
            rows.append(row)
    return rows


def test_registry_workbook_views_follow_repo_reality() -> None:
    repo_root = Path(__file__).resolve().parents[1]
    csv_path = repo_root / "beach_handball_ai/fontes/00_registro/fontes_oficiais.csv"
    xlsx_path = repo_root / "beach_handball_ai/fontes/00_registro/fontes_oficiais.xlsx"
    manifest_path = repo_root / "beach_handball_ai/fontes/05_processado/manifest_checksums.json"
    metadata_path = (
        repo_root
        / "beach_handball_ai/fontes/05_processado/chunks_jsonl/METADADOS_TEMATICOS_ETAPA_3.jsonl"
    )
    stage4_chunks_path = (
        repo_root
        / "beach_handball_ai/fontes/05_processado/chunks_jsonl/CHUNKS_ETAPA_4_CORPUS.jsonl"
    )

    registry_rows = list(csv.DictReader(csv_path.open(encoding="utf-8")))
    registry_by_source = {row["source_id"]: row for row in registry_rows}
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    metadata_records = [
        json.loads(line)
        for line in metadata_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    metadata_by_id = {record["registro_id"]: record for record in metadata_records}
    stage4_records = [
        json.loads(line)
        for line in stage4_chunks_path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    stage4_by_id = {record["chunk_id"]: record for record in stage4_records}

    wb = load_workbook(xlsx_path, read_only=True)

    control_rows = {
        row["source_id"]: row
        for row in registry_rows
        if row["source_id"]
        in {"FONTES_OFICIAIS_CSV", "FONTES_OFICIAIS_XLSX", "MANIFEST_CHECKSUMS_JSON"}
    }
    assert control_rows
    for row in control_rows.values():
        assert row["checksum_sha256"] == "CONTROLE_CICLICO_VER_MANIFESTO"
    manifest_self = next(
        item
        for item in manifest["files"]
        if item["path"] == "beach_handball_ai/fontes/05_processado/manifest_checksums.json"
    )
    assert manifest_self["sha256"] == "AUTOREFERENCIA_EXCLUIDA_DO_HASH"

    audit_rows = sheet_rows(wb["auditoria_organizacao"])
    assert audit_rows
    for row in audit_rows:
        source_id = row["source_id"]
        assert source_id in registry_by_source
        registry_row = registry_by_source[source_id]
        path = repo_root / registry_row["arquivo_local"]
        assert path.exists()
        assert row["arquivo"] == path.name
        assert row["pasta_atual"] == str(path.parent.relative_to(repo_root))
        assert row["pasta_correta"]
        assert row["status_auditoria"]
        assert row["evidencia"]

    action_rows = sheet_rows(wb["proxima_acao"])
    assert len(action_rows) >= 4
    assert any("Etapa 4" in row["acao"] for row in action_rows)
    assert any("G5" in row["criterio_de_conclusao"] or "G5" in row["status"] for row in action_rows)

    review_rows = sheet_rows(wb["revisao_cruzada"])
    assert len(review_rows) == len(metadata_records)
    assert {row["registro_id"] for row in review_rows} == set(metadata_by_id)
    for row in review_rows:
        record = metadata_by_id[row["registro_id"]]
        assert row["source_id"] == record["source_id"]
        assert row["arquivo_base"] == record["arquivo"]
        assert (repo_root / row["arquivo_base"]).exists()
        assert row["status_revisao"] == record["status_pre_chunking"]

    chunk_rows = sheet_rows(wb["matriz_chunks_final"])
    assert len(chunk_rows) == len(stage4_records)
    assert len({row["chunk_id"] for row in chunk_rows}) == len(chunk_rows)
    for row in chunk_rows:
        record = stage4_by_id[row["chunk_id"]]
        assert row["registro_id_origem"] == record["registro_id_origem"]
        assert row["source_id"] == record["source_id"]
        assert row["arquivo_origem"] == record["arquivo_origem"]
        assert (repo_root / row["arquivo_origem"]).exists()
        assert row["status"] == record["status"]
        assert row["uso_no_agente"] == record["uso_no_agente"]
        assert row["acao_necessaria"]
