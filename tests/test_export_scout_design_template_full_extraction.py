import json
from pathlib import Path

from openpyxl import Workbook

from scripts.audit_scout_design_template_full_extraction import audit_extraction
from scripts.export_scout_design_template_full_extraction import build_chunks, write_output


def make_workbook(path: Path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "EVENTOS"
    ws.append(["event_code", "module_id", "status"])
    ws.append(["simple_shot", "finalization_v1", "active_contract"])
    ws.append(["specialist_shot", "legacy", "blocked"])

    module_index = workbook.create_sheet("MODULE_INDEX")
    module_index.append(["module_id", "import_rule_v1", "ui_status"])
    module_index.append(["finalization_v1", "nao_importar_v1", "nao_liberada"])

    fields = workbook.create_sheet("FIELD_DICTIONARY_GLOBAL")
    fields.append(["field_code", "description"])
    fields.append(["scorer_role", "papel dinâmico de pontuação"])

    workbook.save(path)
    workbook.close()


def test_build_chunks_covers_workbook_sheets_and_required_metadata(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "template.xlsx"
    make_workbook(xlsx_path)

    payload = build_chunks(xlsx_path, rows_per_chunk=2)

    assert payload["expected_sheet_count"] == 3
    assert payload["sheet_names"] == ["EVENTOS", "MODULE_INDEX", "FIELD_DICTIONARY_GLOBAL"]
    assert payload["chunk_count"] >= 3

    sheet_names = {chunk["sheet_name"] for chunk in payload["chunks"]}
    assert sheet_names == {"EVENTOS", "MODULE_INDEX", "FIELD_DICTIONARY_GLOBAL"}

    for chunk in payload["chunks"]:
        assert chunk["chunk_id"].startswith("SDT-SHEET-")
        assert chunk["sheet_name"]
        assert chunk["chunk_type"]
        assert chunk["row_range"]
        assert chunk["priority"] in {"critical", "high", "medium", "low"}
        assert "content" in chunk


def test_exported_json_passes_full_extraction_audit(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "full_extraction.json"
    make_workbook(xlsx_path)

    payload = build_chunks(xlsx_path, rows_per_chunk=2)
    write_output(payload, output_path, output_format="json")

    report = audit_extraction(
        chunks_path=output_path,
        xlsx_path=xlsx_path,
        full_extraction=True,
    )

    assert report.ok


def test_exported_jsonl_contains_one_json_object_per_chunk(tmp_path: Path) -> None:
    xlsx_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "full_extraction.jsonl"
    make_workbook(xlsx_path)

    payload = build_chunks(xlsx_path, rows_per_chunk=1)
    write_output(payload, output_path, output_format="jsonl")

    lines = output_path.read_text(encoding="utf-8").splitlines()
    assert len(lines) == payload["chunk_count"]
    first = json.loads(lines[0])
    assert first["chunk_id"] == "SDT-SHEET-0001"
    assert "sheet_name" in first
