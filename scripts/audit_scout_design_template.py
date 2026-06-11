"""Audita o SCOUT_DESIGN_TEMPLATE exportado para XLSX contra o registry v1.

Uso:
    python3 scripts/audit_scout_design_template.py caminho/SCOUT_DESIGN_TEMPLATE.xlsx

Objetivo:
    - verificar se a planilha possui as abas de governanca previstas;
    - validar cabecalhos criticos usados pela IA e por auditoria;
    - comparar eventos/modulos da planilha com scoutpraia.contracts.events_v1;
    - apontar pendencias que ainda impedem declarar o plano 100% implementado.

Este script nao altera a planilha. Ele apenas audita e retorna codigo 1 quando ha
pendencias bloqueantes.
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from scoutpraia.contracts.events_v1 import MODULE_CONTRACTS_V1


REQUIRED_SHEETS = {
    "MODULE_INDEX",
    "SHEET_MAP",
    "EVENTOS",
    "CROSS_MODULE_BOUNDARIES",
    "AI_USE_POLICY",
    "FIELD_DICTIONARY_GLOBAL",
    "RESULT_DOMAIN_GLOBAL",
    "VALIDATION_MATRIX",
    "LEGACY_MIGRATION_RULES",
    "SOURCE_REGISTER",
    "EVENT_REQUIRED_FIELDS",
    "EVENT_OPTIONAL_FIELDS",
    "EVENT_FORBIDDEN_FIELDS",
    "EVENT_BLOCKING_RULES",
    "EVENTOS_LEGADOS_FUTUROS",
    "ARCHITECTURE_README",
}

MODULE_INDEX_REQUIRED_HEADERS = {
    "module_id",
    "primary_events",
    "support_tabs",
    "module_contract_status",
    "import_rule_v1",
    "repo_test_file",
    "evidence_id",
    "ui_status",
    "ai_use_policy",
}

EVENTOS_REQUIRED_HEADERS = {
    "code",
    "module_contract_status",
    "import_rule_v1",
    "active_contract",
    "usable_by_ai",
    "legacy_status",
    "required_result_field",
    "repo_symbol",
}

AI_USE_POLICY_REQUIRED_HEADERS = {
    "allowed_to_suggest",
    "allowed_to_import",
    "requires_human_review",
    "blocked_reason",
    "source_of_truth",
}

SOURCE_REGISTER_REQUIRED_HEADERS = {
    "source_id",
    "organization",
    "version_or_date",
    "module_id",
    "rule_supported",
    "link_or_location",
    "evidence_level",
}

NORMALIZED_RULE_SHEETS = {
    "EVENT_REQUIRED_FIELDS",
    "EVENT_OPTIONAL_FIELDS",
    "EVENT_FORBIDDEN_FIELDS",
    "EVENT_BLOCKING_RULES",
}

LEGACY_CODES = {
    "save",
    "save_shootout",
    "goal_conceded",
    "empty_goal_conceded",
    "fast_break_against",
    "transition_recovery_good",
    "transition_recovery_bad",
    "timeout",
    "set_end",
    "golden_goal",
    "match_end",
    "specialist_shot",
}

MODULE_REPO_TEST_FILES = {
    "finalization_v1": "tests/test_finalization_contract.py",
    "attack_no_shot_v1": "tests/test_attack_no_shot_contract.py",
    "offensive_creation_v1": "tests/test_offensive_creation_contract.py",
    "defensive_v1": "tests/test_defensive_contract.py",
    "shootout_v1": "tests/test_shootout_contract.py",
    "goalkeeper_v1": "tests/test_goalkeeper_contract.py",
    "transition_v1": "tests/test_transition_contract.py",
}


@dataclass
class AuditReport:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warning(self, message: str) -> None:
        self.warnings.append(message)


def _headers(ws) -> list[str]:
    return [str(cell.value).strip() for cell in ws[1] if cell.value is not None]


def _rows_as_dicts(ws) -> list[dict[str, str]]:
    headers = _headers(ws)
    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        values = ["" if value is None else str(value).strip() for value in row[: len(headers)]]
        if not any(values):
            continue
        rows.append(dict(zip(headers, values, strict=False)))
    return rows


def _split_semicolon(value: str) -> set[str]:
    return {item.strip() for item in value.split(";") if item.strip()}


def registry_event_codes() -> set[str]:
    codes: set[str] = set()
    for contract in MODULE_CONTRACTS_V1.values():
        for event in contract.event_contracts():
            codes.add(event.event_code)
    return codes


def registry_primary_event_codes() -> set[str]:
    return {
        event.event_code
        for contract in MODULE_CONTRACTS_V1.values()
        for event in contract.primary_events
    }


def registry_module_codes() -> set[str]:
    return set(MODULE_CONTRACTS_V1)


def validate_required_sheets(wb: Workbook, report: AuditReport) -> None:
    present = set(wb.sheetnames)
    for sheet_name in sorted(REQUIRED_SHEETS - present):
        report.error(f"missing sheet: {sheet_name}")


def validate_headers(wb: Workbook, report: AuditReport) -> None:
    if "MODULE_INDEX" in wb.sheetnames:
        headers = set(_headers(wb["MODULE_INDEX"]))
        for header in sorted(MODULE_INDEX_REQUIRED_HEADERS - headers):
            report.error(f"MODULE_INDEX missing header: {header}")

    if "EVENTOS" in wb.sheetnames:
        headers = set(_headers(wb["EVENTOS"]))
        for header in sorted(EVENTOS_REQUIRED_HEADERS - headers):
            report.error(f"EVENTOS missing header: {header}")

    if "AI_USE_POLICY" in wb.sheetnames:
        headers = set(_headers(wb["AI_USE_POLICY"]))
        for header in sorted(AI_USE_POLICY_REQUIRED_HEADERS - headers):
            report.error(f"AI_USE_POLICY missing header: {header}")

    if "SOURCE_REGISTER" in wb.sheetnames:
        headers = set(_headers(wb["SOURCE_REGISTER"]))
        for header in sorted(SOURCE_REGISTER_REQUIRED_HEADERS - headers):
            report.error(f"SOURCE_REGISTER missing header: {header}")


def validate_module_index(wb: Workbook, report: AuditReport) -> None:
    if "MODULE_INDEX" not in wb.sheetnames:
        return

    rows = _rows_as_dicts(wb["MODULE_INDEX"])
    modules_in_sheet = {row.get("module_id", "") for row in rows}
    for module_code in sorted(registry_module_codes() - modules_in_sheet):
        report.error(f"MODULE_INDEX missing module: {module_code}")

    for row in rows:
        module_id = row.get("module_id", "")
        if not module_id or module_id not in MODULE_CONTRACTS_V1:
            continue
        if row.get("repo_test_file", "") == "":
            report.error(f"MODULE_INDEX row {module_id} missing repo_test_file")
        expected_test = MODULE_REPO_TEST_FILES.get(module_id)
        if expected_test and expected_test not in row.get("repo_test_file", ""):
            report.warning(
                f"MODULE_INDEX row {module_id} repo_test_file does not mention {expected_test}"
            )


def validate_eventos(wb: Workbook, report: AuditReport) -> None:
    if "EVENTOS" not in wb.sheetnames:
        return

    rows = _rows_as_dicts(wb["EVENTOS"])
    codes = {row.get("code", "") for row in rows if row.get("code")}

    for code in sorted(registry_primary_event_codes() - codes):
        report.error(f"EVENTOS missing registry primary event: {code}")

    for code in sorted(LEGACY_CODES - codes):
        # specialist_shot may be intentionally represented only in migration rules.
        report.warning(f"EVENTOS does not list legacy/future code directly: {code}")

    for row in rows:
        code = row.get("code", "")
        if not code:
            continue
        if code in LEGACY_CODES:
            if row.get("usable_by_ai") != "nao_usar_legado_futuro":
                report.error(f"legacy/future code {code} is not blocked by usable_by_ai")
            if row.get("active_contract") != "Não":
                report.error(f"legacy/future code {code} should have active_contract=Não")
        if code in registry_event_codes() and not row.get("repo_symbol"):
            report.error(f"EVENTOS row {code} missing repo_symbol")


def validate_cross_module_boundaries(wb: Workbook, report: AuditReport) -> None:
    if "CROSS_MODULE_BOUNDARIES" not in wb.sheetnames:
        return

    rows = _rows_as_dicts(wb["CROSS_MODULE_BOUNDARIES"])
    pairs = {(row.get("source_module"), row.get("target_module")) for row in rows}
    required_pairs = {
        ("shootout_v1", "goalkeeper_v1"),
        ("transition_v1", "finalization_v1"),
        ("defensive_v1", "goalkeeper_v1"),
    }
    for pair in sorted(required_pairs - pairs):
        report.error(f"CROSS_MODULE_BOUNDARIES missing pair: {pair[0]} -> {pair[1]}")


def validate_legacy_migration(wb: Workbook, report: AuditReport) -> None:
    if "LEGACY_MIGRATION_RULES" not in wb.sheetnames:
        return

    rows = _rows_as_dicts(wb["LEGACY_MIGRATION_RULES"])
    legacy_codes = {row.get("legacy_code", "") for row in rows}
    for code in sorted(LEGACY_CODES - legacy_codes):
        report.error(f"LEGACY_MIGRATION_RULES missing legacy_code: {code}")


def validate_source_register(wb: Workbook, report: AuditReport) -> None:
    if "SOURCE_REGISTER" not in wb.sheetnames:
        return

    rows = _rows_as_dicts(wb["SOURCE_REGISTER"])
    source_ids = {row.get("source_id", "") for row in rows}
    for source_id in ["SRC-001", "SRC-002", "SRC-003", "SRC-004", "SRC-005", "SRC-007", "SRC-008"]:
        if source_id not in source_ids:
            report.error(f"SOURCE_REGISTER missing source_id: {source_id}")


def validate_normalized_rules(wb: Workbook, report: AuditReport) -> None:
    for sheet_name in sorted(NORMALIZED_RULE_SHEETS):
        if sheet_name not in wb.sheetnames:
            continue
        rows = _rows_as_dicts(wb[sheet_name])
        if not rows:
            report.error(f"{sheet_name} has no normalized rules")

    if "VALIDATION_MATRIX" in wb.sheetnames:
        rows = _rows_as_dicts(wb["VALIDATION_MATRIX"])
        val_ids = {row.get("validation_id", "") for row in rows}
        for val_id in ["VAL-013", "VAL-014", "VAL-015"]:
            if val_id not in val_ids:
                report.error(f"VALIDATION_MATRIX missing {val_id}")


def validate_repo_symbol_columns(wb: Workbook, report: AuditReport) -> None:
    for sheet_name in wb.sheetnames:
        if not (
            sheet_name == "EVENTOS"
            or sheet_name.startswith("RESULTADOS_")
            or sheet_name.startswith("CAMPOS_AUXILIARES_")
        ):
            continue
        headers = set(_headers(wb[sheet_name]))
        if "repo_symbol" not in headers:
            report.error(f"{sheet_name} missing repo_symbol header")


def validate_locked_controls(wb: Workbook, report: AuditReport) -> None:
    locked_like_sheets = {"ARCHITECTURE_README", "SOURCE_REGISTER", "VALIDATION_MATRIX"}
    for sheet_name in sorted(locked_like_sheets & set(wb.sheetnames)):
        if not wb[sheet_name].protection.sheet:
            report.warning(f"{sheet_name} is not protected in the XLSX export")


def audit_workbook(path: Path) -> AuditReport:
    wb = load_workbook(path, data_only=True)
    report = AuditReport()

    validate_required_sheets(wb, report)
    validate_headers(wb, report)
    validate_module_index(wb, report)
    validate_eventos(wb, report)
    validate_cross_module_boundaries(wb, report)
    validate_legacy_migration(wb, report)
    validate_source_register(wb, report)
    validate_normalized_rules(wb, report)
    validate_repo_symbol_columns(wb, report)
    validate_locked_controls(wb, report)

    return report


def format_report(report: AuditReport) -> str:
    lines: list[str] = []
    lines.append("== SCOUT_DESIGN_TEMPLATE audit ==")
    lines.append(f"status={'ok' if report.ok else 'failed'}")
    lines.append(f"errors={len(report.errors)}")
    lines.append(f"warnings={len(report.warnings)}")

    if report.errors:
        lines.append("\n-- errors --")
        lines.extend(f"- {error}" for error in report.errors)

    if report.warnings:
        lines.append("\n-- warnings --")
        lines.extend(f"- {warning}" for warning in report.warnings)

    return "\n".join(lines)


def parse_args(argv: Iterable[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Audita SCOUT_DESIGN_TEMPLATE.xlsx contra contracts/events_v1.py")
    parser.add_argument("workbook", type=Path, help="Caminho do XLSX exportado do Google Sheets")
    return parser.parse_args(list(argv))


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    report = audit_workbook(args.workbook)
    print(format_report(report))
    return 0 if report.ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
