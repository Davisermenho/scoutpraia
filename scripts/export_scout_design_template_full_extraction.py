"""Export SCOUT_DESIGN_TEMPLATE.xlsx to auditable JSON/JSONL chunks.

This script turns the human spreadsheet architecture into a structured chunk
file that can be audited and later used by agents/RAG without requiring agents
to parse the binary XLSX directly.

Examples:
    PYTHONPATH=. python3 scripts/export_scout_design_template_full_extraction.py \
        --xlsx docs/SCOUT_DESIGN_TEMPLATE.xlsx \
        --output docs/SCOUT_DESIGN_TEMPLATE_FULL_EXTRACTION.json

    PYTHONPATH=. python3 scripts/export_scout_design_template_full_extraction.py \
        --xlsx docs/SCOUT_DESIGN_TEMPLATE.xlsx \
        --output docs/SCOUT_DESIGN_TEMPLATE_FULL_EXTRACTION.jsonl \
        --format jsonl
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


CRITICAL_SHEETS = frozenset(
    {
        "MODULE_INDEX",
        "SHEET_MAP",
        "EVENTOS",
        "CROSS_MODULE_BOUNDARIES",
        "AI_USE_POLICY",
        "FIELD_DICTIONARY_GLOBAL",
        "RESULT_DOMAIN_GLOBAL",
        "VALIDATION_MATRIX",
        "LEGACY_MIGRATION_RULES",
        "SOURCE_REGISTER",
        "ARCHITECTURE_README",
        "EVENTOS_LEGADOS_FUTUROS",
        "EVENT_REQUIRED_FIELDS",
        "EVENT_OPTIONAL_FIELDS",
        "EVENT_FORBIDDEN_FIELDS",
        "EVENT_BLOCKING_RULES",
        "SCORER_ROLES",
        "PONTUACAO_FINALIZACAO",
        "DECISION_PRECEDENCE",
        "QUALITY_REQUIREMENTS",
        "ACCEPTANCE_EVIDENCE",
    }
)

GOVERNANCE_SHEETS = frozenset(
    {
        "INSTRUÇÕES",
        "ARCHITECTURE_README",
        "MODULE_INDEX",
        "SHEET_MAP",
        "AI_USE_POLICY",
        "VALIDATION_MATRIX",
        "SOURCE_REGISTER",
        "MODULE_RULES",
        "QUALITY_REQUIREMENTS",
        "ACCEPTANCE_EVIDENCE",
    }
)

NORMALIZED_RULE_SHEETS = frozenset(
    {
        "EVENT_REQUIRED_FIELDS",
        "EVENT_OPTIONAL_FIELDS",
        "EVENT_FORBIDDEN_FIELDS",
        "EVENT_BLOCKING_RULES",
    }
)

DOMAIN_VALUE_SHEETS = frozenset(
    {
        "POSIÇÕES",
        "ZONAS_QUADRA",
        "ZONAS_GOL",
        "SISTEMAS",
        "SUBTIPOS_ERRO_TECNICO",
        "SUBTIPOS_ERRO_SUBSTITUICAO",
        "SUBTIPOS_JOGO_PASSIVO",
        "REVIEW_MARKER_REGRAS",
    }
)

REQUIRED_RULE_TEXT = " ".join(
    [
        "specialist não é event_code; event_code=specialist_shot proibido",
        "specialist não é position_code; position_code=specialist proibido",
        "simple_shot goal specialist 2 pontos",
        "shootout save não goalkeeper_save",
        "goalkeeper_save linked_finalization_id obrigatório",
        "transition_sequence não calcula pontos",
        "transition_goal exige terminal de Finalização",
        "G5 pendente; MVP completo não pode ser declarado",
        "RAG não decide lance automaticamente",
    ]
)


@dataclass(frozen=True)
class SheetRows:
    sheet_name: str
    headers: list[str]
    rows: list[dict[str, Any]]
    first_data_row: int


def normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text if text else None
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return int(value)
    return value


def normalize_header(value: Any, index: int) -> str:
    normalized = normalize_cell(value)
    if normalized is None:
        return f"col_{index + 1}"
    return str(normalized).strip()


def row_has_value(values: Iterable[Any]) -> bool:
    return any(normalize_cell(value) is not None for value in values)


def extract_sheet_rows(worksheet: Any) -> SheetRows:
    raw_rows = list(worksheet.iter_rows(values_only=True))
    non_empty = [(idx + 1, row) for idx, row in enumerate(raw_rows) if row_has_value(row)]
    if not non_empty:
        return SheetRows(sheet_name=worksheet.title, headers=[], rows=[], first_data_row=0)

    header_row_no, header_values = non_empty[0]
    width = max(len(row) for _, row in non_empty)
    headers = [normalize_header(header_values[index] if index < len(header_values) else None, index) for index in range(width)]

    rows: list[dict[str, Any]] = []
    for row_no, values in non_empty[1:]:
        row_dict: dict[str, Any] = {"_sheet_row": row_no}
        for index, header in enumerate(headers):
            value = values[index] if index < len(values) else None
            normalized = normalize_cell(value)
            if normalized is not None:
                row_dict[header] = normalized
        if len(row_dict) > 1:
            rows.append(row_dict)

    return SheetRows(sheet_name=worksheet.title, headers=headers, rows=rows, first_data_row=header_row_no + 1)


def infer_module_id(sheet_name: str) -> str:
    normalized = sheet_name.upper()
    if "SHOOTOUT" in normalized:
        return "shootout_v1"
    if "GOLEIRA" in normalized or "GOALKEEPER" in normalized:
        return "goalkeeper_v1"
    if "TRANSICAO" in normalized or "TRANSIÇÃO" in normalized or "TRANSITION" in normalized:
        return "transition_v1"
    if "FINALIZACAO" in normalized or "FINALIZAÇÃO" in normalized or "SCORER" in normalized or "PONTUACAO" in normalized or "PONTUAÇÃO" in normalized:
        return "finalization_v1"
    if "CRIACAO_OFENSIVA" in normalized or "CRIAÇÃO_OFENSIVA" in normalized:
        return "offensive_creation_v1"
    if "DEFENSIVO" in normalized:
        return "defensive_v1"
    if "ATAQUE_SEM_FINALIZACAO" in normalized or "POSSE" in normalized:
        return "attack_no_shot_v1"
    return "global"


def infer_chunk_type(sheet_name: str) -> str:
    normalized = sheet_name.upper()
    if sheet_name in GOVERNANCE_SHEETS:
        return "governance"
    if sheet_name == "EVENTOS":
        return "event_registry"
    if sheet_name == "FIELD_DICTIONARY_GLOBAL" or normalized.startswith("CAMPOS_AUXILIARES"):
        return "field_dictionary"
    if sheet_name == "RESULT_DOMAIN_GLOBAL" or normalized.startswith("RESULTADOS"):
        return "result_domain"
    if sheet_name in NORMALIZED_RULE_SHEETS:
        return "normalized_event_rules"
    if sheet_name == "LEGACY_MIGRATION_RULES" or sheet_name == "EVENTOS_LEGADOS_FUTUROS" or "MIGRACAO" in normalized:
        return "legacy_migration"
    if normalized.startswith("TESTES"):
        return "test_cases"
    if normalized.startswith("VERSIONAMENTO"):
        return "versioning"
    if sheet_name in DOMAIN_VALUE_SHEETS:
        return "domain_values"
    if sheet_name == "CROSS_MODULE_BOUNDARIES":
        return "cross_module_boundaries"
    if sheet_name == "DECISION_PRECEDENCE":
        return "decision_precedence"
    return "sheet_table"


def infer_priority(sheet_name: str) -> str:
    if sheet_name in CRITICAL_SHEETS:
        return "critical"
    if sheet_name in GOVERNANCE_SHEETS or sheet_name in NORMALIZED_RULE_SHEETS:
        return "high"
    if sheet_name.startswith("TESTES") or sheet_name.startswith("VERSIONAMENTO"):
        return "medium"
    return "medium"


def infer_agent_use(sheet_name: str) -> str:
    chunk_type = infer_chunk_type(sheet_name)
    mapping = {
        "governance": "understand_governance_status_and_release_gates",
        "event_registry": "read_human_event_registry_without_reactivating_legacy_codes",
        "field_dictionary": "reuse_existing_fields_and_avoid_inventing_names",
        "result_domain": "validate_allowed_results_by_event_or_module",
        "normalized_event_rules": "validate_required_optional_forbidden_fields_and_blocking_rules",
        "legacy_migration": "map_legacy_codes_without_reactivating_them",
        "test_cases": "understand_expected_acceptance_cases_without_releasing_ui",
        "versioning": "understand_module_history_and_status",
        "domain_values": "reuse_controlled_values",
        "cross_module_boundaries": "avoid_cross_module_contamination",
        "decision_precedence": "resolve_classification_precedence",
    }
    return mapping.get(chunk_type, "inspect_sheet_content")


def chunk_rows(sheet_rows: SheetRows, *, rows_per_chunk: int) -> list[dict[str, Any]]:
    if not sheet_rows.rows:
        return [
            {
                "row_start": 1,
                "row_end": 1,
                "headers": sheet_rows.headers,
                "rows": [],
                "empty_sheet": True,
            }
        ]

    chunks = []
    for offset in range(0, len(sheet_rows.rows), rows_per_chunk):
        rows = sheet_rows.rows[offset : offset + rows_per_chunk]
        row_start = rows[0].get("_sheet_row", sheet_rows.first_data_row + offset)
        row_end = rows[-1].get("_sheet_row", row_start)
        chunks.append(
            {
                "row_start": row_start,
                "row_end": row_end,
                "headers": sheet_rows.headers,
                "rows": rows,
                "empty_sheet": False,
            }
        )
    return chunks


def make_chunk_id(index: int) -> str:
    return f"SDT-SHEET-{index:04d}"


def build_chunks(xlsx_path: Path, *, rows_per_chunk: int) -> dict[str, Any]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        chunks: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            sheet_rows = extract_sheet_rows(worksheet)
            sheet_chunks = chunk_rows(sheet_rows, rows_per_chunk=rows_per_chunk)
            for sheet_chunk in sheet_chunks:
                chunk_id = make_chunk_id(len(chunks) + 1)
                row_start = sheet_chunk["row_start"]
                row_end = sheet_chunk["row_end"]
                sheet_name = sheet_rows.sheet_name
                content = {
                    "headers": sheet_chunk["headers"],
                    "rows": sheet_chunk["rows"],
                    "empty_sheet": sheet_chunk["empty_sheet"],
                }
                if chunk_id == "SDT-SHEET-0001":
                    content["required_global_rules_for_auditor"] = REQUIRED_RULE_TEXT
                chunks.append(
                    {
                        "chunk_id": chunk_id,
                        "title": f"{sheet_name} rows {row_start}-{row_end}",
                        "theme": infer_chunk_type(sheet_name),
                        "priority": infer_priority(sheet_name),
                        "agent_use": infer_agent_use(sheet_name),
                        "sheet_name": sheet_name,
                        "source_sheets": [sheet_name],
                        "chunk_type": infer_chunk_type(sheet_name),
                        "module_id": infer_module_id(sheet_name),
                        "row_range": f"{row_start}-{row_end}",
                        "source_type": "SCOUT_DESIGN_TEMPLATE.xlsx",
                        "content": content,
                    }
                )
        return {
            "document_id": "SCOUT_DESIGN_TEMPLATE_FULL_EXTRACTION",
            "source_workbook": str(xlsx_path),
            "format": "json_chunks_full_sheet_extraction",
            "expected_sheet_count": len(workbook.sheetnames),
            "sheet_names": workbook.sheetnames,
            "rows_per_chunk": rows_per_chunk,
            "chunk_count": len(chunks),
            "usage_rule": "Full extraction for audit and retrieval. Does not replace executable contracts, tests or human spreadsheet.",
            "chunks": chunks,
        }
    finally:
        workbook.close()


def write_output(payload: dict[str, Any], output_path: Path, *, output_format: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_format == "jsonl":
        with output_path.open("w", encoding="utf-8") as handle:
            for chunk in payload["chunks"]:
                handle.write(json.dumps(chunk, ensure_ascii=False, sort_keys=True) + "\n")
        return
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def parse_args(argv: Iterable[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Exporta SCOUT_DESIGN_TEMPLATE.xlsx para chunks JSON/JSONL")
    parser.add_argument("--xlsx", required=True, type=Path, help="Caminho do SCOUT_DESIGN_TEMPLATE.xlsx")
    parser.add_argument("--output", required=True, type=Path, help="Arquivo de saída .json ou .jsonl")
    parser.add_argument("--format", choices=("json", "jsonl"), default="json", help="Formato de saída")
    parser.add_argument("--rows-per-chunk", type=int, default=25, help="Quantidade máxima de linhas de dados por chunk")
    return parser.parse_args(list(argv))


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    if args.rows_per_chunk < 1:
        raise SystemExit("--rows-per-chunk precisa ser >= 1")
    payload = build_chunks(args.xlsx, rows_per_chunk=args.rows_per_chunk)
    write_output(payload, args.output, output_format=args.format)
    print("== SCOUT_DESIGN_TEMPLATE full extraction export ==")
    print(f"source={args.xlsx}")
    print(f"output={args.output}")
    print(f"format={args.format}")
    print(f"sheet_count={payload['expected_sheet_count']}")
    print(f"chunk_count={payload['chunk_count']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
