#!/usr/bin/env python3
"""Sincroniza o registro de fontes RAG: CSV, XLSX, manifest JSON e chunks JSONL.

Escopo:
    - Lê beach_handball_ai/fontes/00_registro/fontes_oficiais.csv
    - Extrai PDFs para Markdown (pdftotext) quando necessário
    - Gera chunks Etapa 4 (CHUNKS_ETAPA_4_CORPUS.jsonl)
    - Atualiza todas as abas do XLSX do registro
    - Atualiza manifest_checksums.json

Modo: MUTANTE — modifica CSV, XLSX, JSONL e Markdown em beach_handball_ai/fontes/
Gate/trigger: após adicionar nova fonte ao registro ou modificar documentos processados
Artefatos produzidos:
    - beach_handball_ai/fontes/00_registro/fontes_oficiais.csv (atualizado)
    - beach_handball_ai/fontes/00_registro/fontes_oficiais.xlsx (atualizado)
    - beach_handball_ai/fontes/05_processado/manifest_checksums.json
    - beach_handball_ai/fontes/05_processado/chunks_jsonl/CHUNKS_ETAPA_4_CORPUS.jsonl
    - beach_handball_ai/fontes/05_processado/chunks_jsonl/CHUNKS_ETAPA_4_AUDITORIA.md

Dependências externas: pdftotext (poppler-utils) para extração de PDF

Examples:
    python3 scripts/refresh_fontes_registry_and_processados.py
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
REGISTRY_CSV = ROOT / "beach_handball_ai/fontes/00_registro/fontes_oficiais.csv"
REGISTRY_XLSX = ROOT / "beach_handball_ai/fontes/00_registro/fontes_oficiais.xlsx"
MANIFEST_JSON = ROOT / "beach_handball_ai/fontes/05_processado/manifest_checksums.json"
PROCESSADOS_DIR = ROOT / "beach_handball_ai/fontes/05_processado/documentos_markdown"
CHUNKS_DIR = ROOT / "beach_handball_ai/fontes/05_processado/chunks_jsonl"
STAGE3_METADATA_JSONL = CHUNKS_DIR / "METADADOS_TEMATICOS_ETAPA_3.jsonl"
STAGE4_CHUNKS_JSONL = CHUNKS_DIR / "CHUNKS_ETAPA_4_CORPUS.jsonl"
STAGE4_CHUNKS_AUDIT_MD = CHUNKS_DIR / "CHUNKS_ETAPA_4_AUDITORIA.md"
LEGACY_IHF_CHUNKS_MD = CHUNKS_DIR / "CHUNKS_FINAIS_RAG_IHF_RULES_BH_2026_PT.md"
TODAY = date.today().isoformat()
CONTROL_ARTIFACT_CHECKSUM = "CONTROLE_CICLICO_VER_MANIFESTO"
MANIFEST_SELF_SHA = "AUTOREFERENCIA_EXCLUIDA_DO_HASH"
CHUNK_MIN_TOKENS = 500
CHUNK_TARGET_TOKENS = 700
CHUNK_MAX_TOKENS = 900
CHUNK_OVERLAP_TOKENS = 100

CATALOG_ROOTS = [ROOT / "docs/sources", ROOT / "beach_handball_ai/fontes"]

FIELDNAMES = [
    "source_id",
    "título",
    "organização",
    "tipo",
    "data",
    "versão",
    "link",
    "uso_permitido",
    "tema",
    "nível_de_confiabilidade",
    "observação",
    "arquivo_local",
    "status",
    "checksum_sha256",
    "data_coleta",
    "source_code_repo",
    "origem_catalogo",
    "papel_documento",
    "fonte_primaria_relacionada",
    "uso_mvp",
    "uso_rag_fase2",
]

MANUAL_REVIEW_OVERRIDES = {
    "CBHB_REGULAMENTO_BH_2026": {
        "status": "revisado_manual_hierarquia_capitulos_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou capítulos e artigos identificáveis no markdown; para tabelas e diagramação específica de competição, preferir o PDF original.",
        "review_note": "Capítulos e artigos principais permaneceram rastreáveis no markdown; tabelas ou blocos de formatação competitiva devem ser conferidos no PDF original.",
    },
    "FHERJ_ESCLARECIMENTOS_REGRAS_2025": {
        "status": "revisado_manual_texto_normativo_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou seções EQUPE, TIRO DE ÁRBITRO, JOGO PASSIVO, SHOOT OUT e PUNIÇÕES legíveis no markdown.",
        "review_note": "Documento predominantemente textual; os blocos normativos centrais ficaram legíveis no markdown gerado.",
    },
    "FHERJ_MUDANCAS_REGRAS_BH_2026": {
        "status": "revisado_manual_texto_normativo_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou os blocos sobre jogo passivo, últimos 15 segundos, proteção do goleiro e equipamento dos jogadores.",
        "review_note": "Mudanças centrais de 2026 ficaram legíveis no markdown, incluindo quatro passes após advertência, últimos 15 segundos e arremesso na cabeça da goleira.",
    },
    "IHF_RULES_BH_2026_EN": {
        "status": "revisado_manual_hierarquia_ok_apendices_identificaveis_lacunas_visuais",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou índice, Rule 7, Clarifications, Substitution Area, Athlete Uniform e Sand Quality/Lighting identificáveis; páginas fotográficas de sinais e layouts visuais exigem o PDF original.",
        "review_note": "A hierarquia normativa e os apêndices críticos ficaram identificáveis. Páginas com sinais ilustrados e layouts visuais de uniforme continuam dependentes do PDF.",
    },
    "IHF_RULES_BH_2026_PT_TRANSLATION": {
        "status": "revisado_manual_hierarquia_ok_apendices_identificaveis_lacunas_visuais",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou índice em português, Regra 7 e apêndices críticos identificáveis; elementos fortemente visuais continuam dependentes do PDF traduzido.",
        "review_note": "A tradução preservou a hierarquia textual e os apêndices críticos. Elementos visuais e tabelas complexas ainda devem ser conferidos no PDF.",
    },
    "IHF_RULES_BH_2026_WORKING_COPY": {
        "status": "revisado_manual_hierarquia_ok_apendices_identificaveis_lacunas_visuais",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou Rule 7 e apêndices críticos identificáveis; sinais ilustrados e componentes de layout visual seguem dependentes do PDF original.",
        "review_note": "A cópia de trabalho preservou a estrutura normativa central. Partes visuais permanecem dependentes do PDF.",
    },
    "MINI_BEACH_HANDBALL_INFO_SHEET": {
        "status": "revisado_manual_texto_ok_box_lateral_inline",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou recomendações de mini/ultimate beach handball legíveis; o box lateral de scoring system foi incorporado em linha no texto extraído.",
        "review_note": "As recomendações principais ficaram legíveis. O quadro lateral de pontuação foi absorvido inline pelo layout extraído.",
    },
    "IHF_RULES_BH_2026_DE": {
        "status": "revisado_manual_hierarquia_ok_apendices_identificaveis_lacunas_visuais",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou índice, Regel 7 e apêndices críticos identificáveis; sinais ilustrados, quadros visuais de uniforme e componentes gráficos seguem dependentes do PDF alemão.",
        "review_note": "A versão alemã preservou a estrutura normativa e os apêndices críticos. Conteúdo visual continua dependente do PDF.",
    },
    "IHF_RULES_BH_2026_FR": {
        "status": "revisado_manual_hierarquia_ok_apendices_identificaveis_lacunas_visuais",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou índice, Règle 7 e apêndices críticos identificáveis; gestos dos árbitros e componentes visuais seguem dependentes do PDF francês.",
        "review_note": "A versão francesa preservou a estrutura normativa e os apêndices críticos. Conteúdo visual continua dependente do PDF.",
    },
    "REFEREEING_BEACH_HANDBALL": {
        "status": "revisado_manual_texto_narrativo_ok_figuras_com_ressalva",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou texto narrativo e marcadores de figuras; gráficos e mapas visuais devem ser consultados no PDF original.",
        "review_note": "O texto técnico ficou legível, mas figuras e gráficos permanecem apenas parcialmente representados no markdown.",
    },
    "SHOOTOUT_PSYCHOLOGICAL_PRESSURE": {
        "status": "revisado_manual_texto_narrativo_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou integridade dos blocos narrativos centrais sobre pressão psicológica, foco e ansiedade competitiva.",
        "review_note": "Documento predominantemente textual; os blocos conceituais principais permaneceram legíveis no markdown.",
    },
    "SRC_IHF_RULES_FILE": {
        "status": "revisado_manual_hierarquia_ok_apendices_identificaveis_lacunas_visuais",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida equivalente à cópia operacional IHF em inglês: Rule 7 e apêndices identificáveis; conteúdo fotográfico e layouts visuais exigem o PDF original.",
        "review_note": "A fonte canônica do repositório preservou a estrutura textual principal; componentes visuais seguem dependentes do PDF.",
    },
    "SRC_NOTATIONAL_BH_IANNACCONE_2022": {
        "status": "revisado_manual_artigo_tabelas_principais_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou Abstract, Introduction, Methods, Results, Conclusions e Tabelas 1–6 legíveis no markdown processado.",
        "review_note": "O artigo preservou resumo, método, resultados e tabelas principais com legibilidade suficiente para chunking textual; a figura esquemática continua melhor no PDF.",
    },
    "SRC_RAG_STRUCTURED_FILE": {
        "status": "revisado_manual_artigo_tabelas_principais_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou Abstract, Introduction, Methodology, Results, Conclusion e Tabelas 1–7 identificáveis no markdown processado.",
        "review_note": "O artigo preservou as seções centrais e as tabelas principais; diagramas e exemplos visuais seguem melhor representados no PDF.",
    },
    "SRC_WOMENS_BH_STATISTICS_2022": {
        "status": "revisado_manual_artigo_tabelas_principais_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou Abstract, Introduction, Methods, Results, Discussion and conclusions e Tabelas 1–4 legíveis no markdown processado.",
        "review_note": "O artigo preservou as seções centrais e as tabelas principais com legibilidade suficiente para uso textual; formatação fina permanece melhor no PDF.",
    },
    "ULTIMATE_SCHOOL_HANDBALL_2025": {
        "status": "revisado_manual_texto_narrativo_ok",
        "uso_rag_fase2": "revisado_manual_bloqueado_ate_fase2_global",
        "observação": "Revisão manual assistida confirmou a introdução e os blocos pedagógicos centrais em texto limpo suficiente para leitura técnica.",
        "review_note": "Documento predominantemente textual; a estrutura argumentativa principal permaneceu legível no markdown.",
    },
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_registry() -> list[dict[str, str]]:
    with REGISTRY_CSV.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def write_registry(rows: list[dict[str, str]]) -> None:
    rows = sorted(rows, key=lambda row: row["source_id"])
    with REGISTRY_CSV.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def normalize_row(row: dict[str, str]) -> dict[str, str]:
    normalized = {field: (row.get(field, "") or "") for field in FIELDNAMES}
    path = ROOT / normalized["arquivo_local"]
    if path.exists():
        if is_cyclic_control_artifact(path):
            normalized["checksum_sha256"] = CONTROL_ARTIFACT_CHECKSUM
        else:
            normalized["checksum_sha256"] = sha256(path)
    if not normalized["data_coleta"]:
        normalized["data_coleta"] = TODAY
    if not normalized["origem_catalogo"]:
        normalized["origem_catalogo"] = infer_catalog_origin(path)
    if not normalized["data"] and not normalized["versão"]:
        normalized["versão"] = infer_version(normalized)
    return normalized


def is_cyclic_control_artifact(path: Path) -> bool:
    return path in {REGISTRY_CSV, REGISTRY_XLSX, MANIFEST_JSON}


def infer_catalog_origin(path: Path) -> str:
    if path.is_relative_to(ROOT / "docs/sources"):
        return "docs/sources"
    return "beach_handball_ai/fontes"


def infer_version(row: dict[str, str]) -> str:
    path = Path(row["arquivo_local"])
    if "04_fontes_proprias_cepraea" in row["arquivo_local"]:
        return "v0"
    if "05_processado" in row["arquivo_local"]:
        return TODAY
    if path.name == "README.md":
        return "v1"
    if path.name == "Working-with-evals.md":
        return "snapshot_2026-06-11"
    if path.name == "Scout de Handebol de Areia_ Fontes Fortes.md":
        return "curadoria_v1"
    return TODAY


def load_stage3_metadata() -> list[dict[str, object]]:
    if not STAGE3_METADATA_JSONL.exists():
        return []

    records = []
    for line in STAGE3_METADATA_JSONL.read_text(encoding="utf-8").splitlines():
        if line.strip():
            records.append(json.loads(line))
    return records


def load_stage4_chunks() -> list[dict[str, object]]:
    if not STAGE4_CHUNKS_JSONL.exists():
        return []

    records = []
    for line in STAGE4_CHUNKS_JSONL.read_text(encoding="utf-8").splitlines():
        if line.strip():
            records.append(json.loads(line))
    return records


def token_count(text: str) -> int:
    return len(re.findall(r"\S+", text))


def slugify(text: str) -> str:
    normalized = text.lower()
    replacements = {
        "á": "a",
        "à": "a",
        "â": "a",
        "ã": "a",
        "é": "e",
        "ê": "e",
        "í": "i",
        "ó": "o",
        "ô": "o",
        "õ": "o",
        "ú": "u",
        "ç": "c",
    }
    for source, target in replacements.items():
        normalized = normalized.replace(source, target)
    normalized = re.sub(r"[^a-z0-9]+", "_", normalized)
    return normalized.strip("_")


def strip_frontmatter(text: str) -> str:
    if not text.startswith("---\n"):
        return text
    marker = "\n---\n"
    end = text.find(marker, 4)
    if end == -1:
        return text
    return text[end + len(marker) :].lstrip()


def extract_title_line(text: str, fallback: str) -> str:
    for line in strip_frontmatter(text).splitlines():
        cleaned = line.strip()
        if cleaned.startswith("#"):
            return cleaned.lstrip("#").strip()
    return fallback


def extract_overlap_text(text: str, overlap_tokens: int) -> str:
    words = re.findall(r"\S+", text)
    if len(words) <= overlap_tokens:
        return text
    return " ".join(words[-overlap_tokens:])


def split_long_paragraph(paragraph: str, target_tokens: int) -> list[str]:
    if token_count(paragraph) <= CHUNK_MAX_TOKENS:
        return [paragraph]

    sentences = re.split(r"(?<=[.!?])\s+", paragraph.strip())
    chunks = []
    current = []
    current_tokens = 0
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        sentence_tokens = token_count(sentence)
        if current and current_tokens + sentence_tokens > target_tokens:
            chunks.append(" ".join(current).strip())
            current = [sentence]
            current_tokens = sentence_tokens
        else:
            current.append(sentence)
            current_tokens += sentence_tokens
    if current:
        chunks.append(" ".join(current).strip())
    return [chunk for chunk in chunks if chunk]


def split_into_paragraphs(text: str) -> list[str]:
    cleaned = re.sub(r"\n{3,}", "\n\n", text.strip())
    paragraphs = []
    for block in re.split(r"\n\s*\n", cleaned):
        normalized = block.strip()
        if not normalized:
            continue
        paragraphs.extend(split_long_paragraph(normalized, CHUNK_TARGET_TOKENS))
    return paragraphs


def extract_anchor_slice(text: str, anchor: str, sibling_anchors: list[str]) -> str:
    if not anchor:
        return text
    start = text.find(anchor)
    if start == -1:
        return text

    end = len(text)
    for sibling in sibling_anchors:
        if sibling == anchor:
            continue
        sibling_start = text.find(sibling, start + len(anchor))
        if sibling_start != -1:
            end = min(end, sibling_start)

    heading_start = text.rfind("\n#", 0, start)
    if heading_start != -1:
        start = heading_start + 1
    return text[start:end].strip()


def source_row_index(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {row["source_id"]: row for row in rows}


def infer_chunk_status(record: dict[str, object], source_rows: dict[str, dict[str, str]]) -> str:
    source_row = source_rows.get(str(record["source_id"]))
    source_status = (source_row or {}).get("status", "")
    if "_rascunhos/" in str(record["arquivo"]):
        return "deprecated"
    if "lacunas_visuais" in source_status or "figuras_com_ressalva" in source_status:
        return "ativo_com_ressalva"
    if "duplicata_linguistica" in source_status:
        return "ativo_com_ressalva"
    return "ativo"


def infer_chunk_usage(record: dict[str, object], status: str) -> str:
    if status == "deprecated":
        return "nao_usar_em_ingestao"
    if status.startswith("bloqueado"):
        return "responder_nao_encontrado_ou_consultar_pdf"
    return chunking_usage(record)


def build_chunk_text(record: dict[str, object], title: str, body: str) -> str:
    anchor = str(record.get("ancora", "")).strip()
    parts = [f"# {title}"]
    if anchor and anchor not in title and anchor in body:
        parts.extend(["", f"Trecho-guia: {anchor}"])
    parts.extend(["", body.strip()])
    return "\n".join(parts).strip()


def window_paragraph_chunks(base_text: str, chunk_title: str) -> list[tuple[str, int]]:
    paragraphs = split_into_paragraphs(base_text)
    if not paragraphs:
        return []

    chunks = []
    current: list[str] = []
    current_tokens = 0
    previous_chunk_text = ""
    overlap_tokens = 0

    for paragraph in paragraphs:
        paragraph_tokens = token_count(paragraph)
        if current and current_tokens >= CHUNK_MIN_TOKENS and current_tokens + paragraph_tokens > CHUNK_MAX_TOKENS:
            chunk_body = "\n\n".join(current).strip()
            chunk_text = "\n".join([chunk_title, "", chunk_body]).strip()
            chunks.append((chunk_text, overlap_tokens))
            previous_chunk_text = chunk_body
            overlap_excerpt = extract_overlap_text(previous_chunk_text, CHUNK_OVERLAP_TOKENS)
            current = [overlap_excerpt, paragraph]
            current_tokens = token_count(overlap_excerpt) + paragraph_tokens
            overlap_tokens = min(CHUNK_OVERLAP_TOKENS, token_count(overlap_excerpt))
        else:
            current.append(paragraph)
            current_tokens += paragraph_tokens

    if current:
        chunk_body = "\n\n".join(current).strip()
        chunk_text = "\n".join([chunk_title, "", chunk_body]).strip()
        chunks.append((chunk_text, overlap_tokens))

    return chunks


def parse_legacy_ihf_chunks() -> tuple[list[dict[str, object]], dict[str, list[str]]]:
    if not LEGACY_IHF_CHUNKS_MD.exists():
        return [], {}

    text = LEGACY_IHF_CHUNKS_MD.read_text(encoding="utf-8")
    sections = re.split(r"\n## ", text)
    records = []
    replacement_map: dict[str, list[str]] = {}
    part_to_registro = {
        "P01": "ihf_rules_revision_part_01",
        "P02": "ihf_rules_revision_part_02",
        "P03": "ihf_rules_revision_part_03",
        "P04": "ihf_rules_revision_part_04",
        "P05": "ihf_rules_revision_part_05",
        "P06": "ihf_rules_revision_part_06",
    }

    for section in sections:
        section = section.strip()
        if not section.startswith("CHUNK_"):
            continue
        lines = section.splitlines()
        heading = lines[0].strip()
        match = re.match(r"^(CHUNK_[A-Z0-9_]+)\s+—\s+(.*)$", heading)
        if not match:
            continue
        chunk_id, title = match.groups()
        status_match = re.search(r"\nstatus:\s*(.+)", "\n" + section)
        status_raw = status_match.group(1).strip() if status_match else "liberado"
        if "BLOQUEADO VISUAL" in title.upper():
            status = "bloqueado_visual"
        elif status_raw == "liberado":
            status = "ativo"
        elif status_raw == "liberado_com_ressalva":
            status = "ativo_com_ressalva"
        elif status_raw == "bloqueado":
            status = "bloqueado"
        else:
            status = slugify(status_raw)

        uso_match = re.search(r"\nuso_no_agente:\s*(.+)", "\n" + section)
        uso_no_agente = uso_match.group(1).strip() if uso_match else "normativo_prioritario"
        ressalva_match = re.search(r"\nressalva:\s*(.+)", "\n" + section)
        ressalva = ressalva_match.group(1).strip() if ressalva_match else ""
        content_match = re.search(r"\nConteúdo:\s*(.+)", "\n" + section, re.DOTALL)
        decision_match = re.search(r"\ndecisao:\s*(.+)", "\n" + section, re.DOTALL)
        body = ""
        if content_match:
            body = content_match.group(1).strip()
        elif decision_match:
            body = decision_match.group(1).strip()

        part_match = re.search(r"CHUNK_(P\d{2})", chunk_id)
        part_code = part_match.group(1) if part_match else "P00"
        registro_id = part_to_registro.get(part_code, "ihf_legacy_chunks")
        replacement_map.setdefault(registro_id, []).append(chunk_id)

        tema = "arbitragem" if part_code == "P05" else "regra"
        record = {
            "chunk_id": chunk_id,
            "registro_id_origem": registro_id,
            "source_id": "IHF_RULES_BH_2026_EN",
            "source_id_operacional": "IHF_RULES_BH_2026_PT_TRANSLATION",
            "tema": tema,
            "subtema": slugify(title),
            "organizacao": "IHF",
            "versao": "2026",
            "confiabilidade": "A",
            "arquivo_origem": str(LEGACY_IHF_CHUNKS_MD.relative_to(ROOT)),
            "status": status,
            "metodo_chunking": "curadoria_legada_convertida",
            "uso_no_agente": uso_no_agente,
            "overlap_tokens_aprox": 0,
            "token_count_aprox": token_count(body),
            "observacao": ressalva or "Chunk curado legado convertido para estrutura JSONL da Etapa 4.",
            "texto": body,
        }
        records.append(record)

    return records, replacement_map


def build_deprecated_stage3_chunk_records(
    metadata_records: list[dict[str, object]],
    replacement_map: dict[str, list[str]],
) -> list[dict[str, object]]:
    deprecated = []
    for record in metadata_records:
        registro_id = str(record["registro_id"])
        if not registro_id.startswith("ihf_rules_revision_part_"):
            continue
        replacements = replacement_map.get(registro_id, [])
        deprecated_text = (
            "Rascunho historico substituido na Etapa 4. "
            "Nao ingerir no banco principal; usar apenas para auditoria do processo."
        )
        deprecated.append(
            {
                "chunk_id": f"DEPRECATED_{registro_id.upper()}",
                "registro_id_origem": registro_id,
                "source_id": str(record["source_id"]),
                "source_id_operacional": str(record.get("source_id_operacional", "")),
                "tema": str(record["tema"]),
                "subtema": str(record["subtema"]),
                "organizacao": str(record["organizacao"]),
                "versao": str(record["versao"]),
                "confiabilidade": str(record["confiabilidade"]),
                "arquivo_origem": str(record["arquivo"]),
                "status": "deprecated",
                "metodo_chunking": "placeholder_deprecated",
                "uso_no_agente": "nao_usar_em_ingestao",
                "overlap_tokens_aprox": 0,
                "token_count_aprox": token_count(deprecated_text),
                "observacao": (
                    "Rascunho historico substituido pelos chunks estruturados "
                    + ", ".join(replacements)
                    if replacements
                    else "Rascunho historico mantido apenas para auditoria do processo."
                ),
                "replacement_chunk_ids": replacements,
                "texto": deprecated_text,
            }
        )
    return deprecated


def build_source_text_map(metadata_records: list[dict[str, object]]) -> dict[str, str]:
    source_text = {}
    for record in metadata_records:
        arquivo = str(record["arquivo"])
        if arquivo in source_text:
            continue
        source_text[arquivo] = strip_frontmatter((ROOT / arquivo).read_text(encoding="utf-8"))
    return source_text


def build_stage4_generated_chunks(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    metadata_records = load_stage3_metadata()
    source_rows = source_row_index(rows)
    source_text = build_source_text_map(metadata_records)
    grouped_by_file: dict[str, list[dict[str, object]]] = {}
    for record in metadata_records:
        grouped_by_file.setdefault(str(record["arquivo"]), []).append(record)

    generated: list[dict[str, object]] = []
    chunk_sequence = 1

    for record in metadata_records:
        registro_id = str(record["registro_id"])
        if registro_id.startswith("ihf_rules_revision_part_"):
            continue

        arquivo = str(record["arquivo"])
        body = source_text[arquivo]
        siblings = grouped_by_file[arquivo]
        sibling_anchors = [str(sibling.get("ancora", "")).strip() for sibling in siblings if sibling.get("ancora")]
        body = extract_anchor_slice(body, str(record.get("ancora", "")).strip(), sibling_anchors)
        title = extract_title_line(body, registro_id)
        status = infer_chunk_status(record, source_rows)
        base_chunk_title = f"# {title}"
        chunk_windows = window_paragraph_chunks(body, base_chunk_title)
        if not chunk_windows:
            chunk_windows = [(build_chunk_text(record, title, body), 0)]

        for index, (chunk_text, overlap_tokens) in enumerate(chunk_windows, start=1):
            generated.append(
                {
                    "chunk_id": f"ETAPA4_{chunk_sequence:03d}_{slugify(registro_id)[:36]}",
                    "registro_id_origem": registro_id,
                    "source_id": str(record["source_id"]),
                    "source_id_operacional": str(record.get("source_id_operacional", "")),
                    "tema": str(record["tema"]),
                    "subtema": str(record["subtema"]),
                    "organizacao": str(record["organizacao"]),
                    "versao": str(record["versao"]),
                    "confiabilidade": str(record["confiabilidade"]),
                    "arquivo_origem": arquivo,
                    "status": status,
                    "metodo_chunking": (
                        "janela_unica_sem_split" if len(chunk_windows) == 1 else "janela_paragrafos_com_overlap"
                    ),
                    "uso_no_agente": infer_chunk_usage(record, status),
                    "overlap_tokens_aprox": overlap_tokens,
                    "token_count_aprox": token_count(chunk_text),
                    "observacao": (
                        str(record.get("observacao", ""))
                        if len(chunk_windows) == 1
                        else f"{record.get('observacao', '')} | parte {index}/{len(chunk_windows)}"
                    ),
                    "texto": chunk_text.strip(),
                }
            )
            chunk_sequence += 1

    return generated


def sync_stage4_chunks(rows: list[dict[str, str]]) -> None:
    metadata_records = load_stage3_metadata()
    legacy_records, replacement_map = parse_legacy_ihf_chunks()
    generated_records = build_stage4_generated_chunks(rows)
    deprecated_records = build_deprecated_stage3_chunk_records(metadata_records, replacement_map)

    all_records = sorted(
        legacy_records + generated_records + deprecated_records,
        key=lambda record: (str(record["status"]) == "deprecated", str(record["chunk_id"])),
    )
    STAGE4_CHUNKS_JSONL.write_text(
        "\n".join(json.dumps(record, ensure_ascii=False) for record in all_records) + "\n",
        encoding="utf-8",
    )

    status_counts = Counter(str(record["status"]) for record in all_records)
    tema_counts = Counter(str(record["tema"]) for record in all_records)
    lines = [
        "# Auditoria Etapa 4",
        "",
        f"gerado_em: {TODAY}",
        f"arquivo_jsonl: {STAGE4_CHUNKS_JSONL.relative_to(ROOT)}",
        f"chunks_totais: {len(all_records)}",
        "",
        "## Status",
        "",
    ]
    for status, count in sorted(status_counts.items()):
        lines.append(f"- {status}: {count}")
    lines.extend(["", "## Temas", ""])
    for tema, count in sorted(tema_counts.items()):
        lines.append(f"- {tema}: {count}")
    lines.extend(
        [
            "",
            "## Regras operacionais",
            "",
            "- chunks derivados de `_rascunhos/` ficam com `status=deprecated` e nao entram na ingestao principal.",
            "- chunks normativos IHF convertidos do artefato legado continuam separados de CEPRAEA por `tema`, `organizacao` e `source_id`.",
            "- a Etapa 4 continua documental; embeddings/Chroma seguem bloqueados ate o gate global do ScoutPraia.",
            "",
        ]
    )
    STAGE4_CHUNKS_AUDIT_MD.write_text("\n".join(lines), encoding="utf-8")


def expected_operational_folder(row: dict[str, str]) -> str:
    path = ROOT / row["arquivo_local"]
    org = row["organização"]

    if "04_fontes_proprias_cepraea" in row["arquivo_local"]:
        return "beach_handball_ai/fontes/04_fontes_proprias_cepraea"
    if org in {"IHF", "FHERJ", "Tradução técnica - Luiz Filipe Galvão da Silva Caldas"}:
        return "beach_handball_ai/fontes/01_ihf_regras"
    if org == "EHF":
        return "beach_handball_ai/fontes/02_ehf_tecnico"
    if org == "CBHb":
        return "beach_handball_ai/fontes/03_cbhb_brasil"
    return str(path.parent.relative_to(ROOT))


def operational_source_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    filtered = []
    for row in rows:
        path = row["arquivo_local"]
        if not path.startswith("beach_handball_ai/fontes/0"):
            continue
        if "/00_registro/" in path or "/05_processado/" in path:
            continue
        filtered.append(row)
    return sorted(filtered, key=lambda row: row["source_id"])


def chunking_usage(record: dict[str, object]) -> str:
    source_id = str(record["source_id"])
    organizacao = str(record["organizacao"])
    tema = str(record["tema"])

    if organizacao == "IHF":
        return "normativo_prioritario" if tema in {"regra", "arbitragem"} else "apoio_oficial"
    if organizacao == "CBHb":
        return "competicao_brasil"
    if organizacao == "EHF":
        return "apoio_tecnico_educacional"
    if source_id.startswith("CEPRAEA_") or organizacao == "CEPRAEA":
        return "nomenclatura_e_metodo_interno"
    return "apoio_controlado"


def chunking_risk(record: dict[str, object]) -> str:
    arquivo = str(record["arquivo"])
    organizacao = str(record["organizacao"])
    tema = str(record["tema"])

    if "_rascunhos/" in arquivo:
        return "Base historica revisada; usar apenas como insumo para chunk final rastreavel."
    if organizacao == "EHF":
        return "Nao substituir regra oficial IHF quando a pergunta for normativa."
    if organizacao == "CBHb":
        return "Aplicacao contextual brasileira; nao expandir para regra universal."
    if organizacao == "CEPRAEA":
        return "Nao misturar convencao interna do CEPRAEA com regra oficial no mesmo chunk."
    if tema in {"regra", "arbitragem"}:
        return "Manter precedencia normativa e separar atualizacao oficial de comentario interno."
    return "Sem risco estrutural adicional alem da rastreabilidade obrigatoria."


def build_auditoria_organizacao_rows(rows: list[dict[str, str]]) -> tuple[list[str], list[list[str]]]:
    header = [
        "source_id",
        "arquivo",
        "pasta_atual",
        "pasta_correta",
        "conteudo_lido",
        "nivel",
        "status_auditoria",
        "problema_real",
        "proxima_acao",
        "evidencia",
    ]
    records = []
    for row in operational_source_rows(rows):
        path = ROOT / row["arquivo_local"]
        processed_md = (
            ROOT
            / "beach_handball_ai/fontes/05_processado/documentos_markdown"
            / f"MD_PROCESSADO__{row['source_id']}.md"
        )
        pasta_atual = str(path.parent.relative_to(ROOT))
        pasta_correta = expected_operational_folder(row)
        conteudo_lido = "Sim" if path.suffix.lower() != ".pdf" or processed_md.exists() else "Nao"
        if pasta_atual != pasta_correta:
            status_auditoria = "fora_do_padroes"
            problema = "Arquivo registrado fora da pasta esperada pelo corpus operacional."
            proxima_acao = "Mover ou corrigir o registro antes de qualquer chunking."
        elif "lacunas_visuais" in row["status"] or "figuras_com_ressalva" in row["status"]:
            status_auditoria = "coerente_com_ressalva_visual"
            problema = "Elementos visuais, tabelas ou figuras ainda dependem do PDF original."
            proxima_acao = "Manter no diretorio atual e consultar o PDF quando a pergunta depender de visual."
        elif "duplicata_linguistica" in row["status"]:
            status_auditoria = "coerente_duplicata_de_conferencia"
            problema = "Duplicata linguistica de apoio; nao e a fonte operacional principal."
            proxima_acao = "Usar apenas para conferencia linguistica, preservando EN/PT como referencia operacional."
        elif row["organização"] == "CBHb":
            status_auditoria = "coerente_origem_privada"
            problema = "Origem operacional por WhatsApp; nao ha link publico oficial rastreavel no repo."
            proxima_acao = "Manter a ressalva de origem privada no registro e usar apenas no contexto brasileiro."
        elif row["organização"] == "CEPRAEA":
            status_auditoria = "coerente_base_interna"
            problema = "Convencao interna; nao pode ser promovida a regra oficial."
            proxima_acao = "Separar dos chunks normativos e manter precedencia abaixo de IHF/CBHb/EHF."
        else:
            status_auditoria = "coerente_com_repo"
            problema = "Nenhum desalinhamento estrutural identificado."
            proxima_acao = "Manter o registro e a pasta atuais."

        evidencia = (
            f"{row['source_id']} | {row['arquivo_local']}"
            if path.suffix.lower() != ".pdf"
            else f"{row['source_id']} | {processed_md.relative_to(ROOT)}"
        )
        records.append(
            [
                row["source_id"],
                path.name,
                pasta_atual,
                pasta_correta,
                conteudo_lido,
                row["nível_de_confiabilidade"],
                status_auditoria,
                problema,
                proxima_acao,
                evidencia,
            ]
        )
    return header, records


def build_proxima_acao_rows() -> tuple[list[str], list[list[str]]]:
    header = ["ordem", "acao", "criterio_de_conclusao", "responsavel", "status"]
    if STAGE4_CHUNKS_JSONL.exists():
        rows = [
            [
                "1",
                "Tratar checksums ciclicos dos artefatos de controle (CSV/XLSX/manifest) sem autoreferencia.",
                "Registro e workbook exibem sentinela controlada; manifest preserva hash real de CSV/XLSX e nao tenta hashear a si mesmo.",
                "Execucao tecnica",
                "concluido",
            ],
            [
                "2",
                "Manter auditoria_organizacao, revisao_cruzada e matriz_chunks_final como visoes derivadas do estado canonico atual.",
                "As abas nao apontam para arquivos removidos ou caminhos inexistentes do repo.",
                "Execucao tecnica",
                "concluido",
            ],
            [
                "3",
                "Etapa 4 executada: chunks reais do corpus gerados com status por unidade e historicos deprecated.",
                "CHUNKS_ETAPA_4_CORPUS.jsonl cobre o inventario aprovado da Etapa 3 com source_id, tema, confiabilidade e separacao CEPRAEA x regra oficial.",
                "Execucao tecnica",
                "concluido_neste_ciclo",
            ],
            [
                "4",
                "Etapa 5 permanece bloqueada: nao gerar embeddings nem Chroma antes do gate global do ScoutPraia.",
                "G5 aprovado e criterios de docs/009_RAG_Workflow_Fontes.md satisfeitos antes das Etapas 5-8.",
                "Governanca do projeto",
                "bloqueado_ate_G5",
            ],
        ]
    else:
        rows = [
            [
                "1",
                "Tratar checksums ciclicos dos artefatos de controle (CSV/XLSX/manifest) sem autoreferencia.",
                "Registro e workbook exibem sentinela controlada; manifest preserva hash real de CSV/XLSX e nao tenta hashear a si mesmo.",
                "Execucao tecnica",
                "concluido_neste_ciclo",
            ],
            [
                "2",
                "Manter auditoria_organizacao, revisao_cruzada e matriz_chunks_final como visoes derivadas do estado canonico atual.",
                "As abas nao apontam para arquivos removidos ou caminhos inexistentes do repo.",
                "Execucao tecnica",
                "concluido_neste_ciclo",
            ],
            [
                "3",
                "Executar a Etapa 4 do plano: gerar chunks reais de 500-900 tokens com sobreposicao controlada e source_id obrigatorio.",
                "Cada chunk fica compreensivel sozinho, separado por precedencia (IHF/CBHb/EHF/CEPRAEA) e marcado como deprecated quando historico.",
                "Proxima execucao",
                "acao_recomendada_imediata",
            ],
            [
                "4",
                "Nao iniciar embeddings, Chroma ou avaliacao do agente textual antes do gate global do ScoutPraia.",
                "G5 aprovado e criterios de docs/009_RAG_Workflow_Fontes.md satisfeitos antes das Etapas 5-8.",
                "Governanca do projeto",
                "bloqueado_ate_G5",
            ],
        ]
    return header, rows


def build_revisao_cruzada_rows() -> tuple[list[str], list[list[str]]]:
    header = [
        "registro_id",
        "source_id",
        "source_id_operacional",
        "arquivo_base",
        "tema",
        "subtema",
        "status_revisao",
        "risco_identificado",
        "criterio_para_chunking",
        "observacao",
    ]
    rows = []
    for record in load_stage3_metadata():
        pronto = bool(record["pronto_para_chunking"])
        criterio = (
            "Pode entrar na Etapa 4 mantendo source_id, tema, subtema e precedencia de fonte."
            if pronto
            else "Nao chunkar enquanto a revisao manual da Etapa 2 nao estiver fechada."
        )
        rows.append(
            [
                str(record["registro_id"]),
                str(record["source_id"]),
                str(record.get("source_id_operacional", "")),
                str(record["arquivo"]),
                str(record["tema"]),
                str(record["subtema"]),
                str(record["status_pre_chunking"]),
                chunking_risk(record),
                criterio,
                str(record.get("observacao", "")),
            ]
        )
    return header, rows


def build_matriz_chunks_rows() -> tuple[list[str], list[list[str]]]:
    stage4_chunks = load_stage4_chunks()
    if stage4_chunks:
        header = [
            "chunk_id",
            "registro_id_origem",
            "source_id",
            "tema",
            "subtema",
            "status",
            "arquivo_origem",
            "uso_no_agente",
            "acao_necessaria",
        ]
        rows = []
        for chunk in stage4_chunks:
            rows.append(
                [
                    str(chunk["chunk_id"]),
                    str(chunk["registro_id_origem"]),
                    str(chunk["source_id"]),
                    str(chunk["tema"]),
                    str(chunk["subtema"]),
                    str(chunk["status"]),
                    str(chunk["arquivo_origem"]),
                    str(chunk["uso_no_agente"]),
                    (
                        "Nao ingerir; mantido apenas para auditoria."
                        if str(chunk["status"]) == "deprecated"
                        else "Pode sustentar Etapa 5 apenas quando o gate global do ScoutPraia liberar RAG."
                    ),
                ]
            )
        return header, rows

    header = [
        "chunk_planejado_id",
        "registro_id_origem",
        "source_id",
        "tema",
        "subtema",
        "status_atual",
        "arquivo_base",
        "uso_no_agente",
        "acao_necessaria",
    ]
    rows = []
    for index, record in enumerate(load_stage3_metadata(), start=1):
        pronto = bool(record["pronto_para_chunking"])
        rows.append(
            [
                f"CHUNK_PLAN_{index:03d}",
                str(record["registro_id"]),
                str(record["source_id"]),
                str(record["tema"]),
                str(record["subtema"]),
                "pronto_para_etapa_4" if pronto else "bloqueado_pre_chunking",
                str(record["arquivo"]),
                chunking_usage(record),
                (
                    "Gerar chunk real sem misturar precedencia normativa e marcar historicos como deprecated."
                    if pronto
                    else "Concluir revisao antes de qualquer chunking."
                ),
            ]
        )
    return header, rows


def base_processed_row(source_row: dict[str, str], output_rel: str) -> dict[str, str]:
    source_id = source_row["source_id"]
    row = {
        "source_id": f"MD_PROCESSADO_{source_id}",
        "título": f"MD_PROCESSADO__{source_id}",
        "organização": "ScoutPraia/CEPRAEA",
        "tipo": "Artefato processado markdown",
        "data": TODAY,
        "versão": TODAY,
        "link": "",
        "uso_permitido": "uso interno operacional",
        "tema": "processamento/intermediario",
        "nível_de_confiabilidade": "derivado",
        "observação": f"Conversão local por pdftotext do PDF {source_id}, mantendo rastreabilidade explícita ao source_id original.",
        "arquivo_local": output_rel,
        "status": "texto_extraido_layout_pendente_revisao_manual_tabelas",
        "checksum_sha256": "",
        "data_coleta": TODAY,
        "source_code_repo": source_row.get("source_code_repo", ""),
        "origem_catalogo": "beach_handball_ai/fontes",
        "papel_documento": "artefato_processado_consolidado",
        "fonte_primaria_relacionada": source_id,
        "uso_mvp": "nao",
        "uso_rag_fase2": "bloqueado_ate_revisao_manual",
    }
    override = MANUAL_REVIEW_OVERRIDES.get(source_id)
    if override:
        row["status"] = override["status"]
        row["uso_rag_fase2"] = override["uso_rag_fase2"]
        row["observação"] = override["observação"]
    return row


def extract_pdf_to_markdown(source_row: dict[str, str]) -> tuple[str, str]:
    pdf_path = ROOT / source_row["arquivo_local"]
    output_name = f"MD_PROCESSADO__{source_row['source_id']}.md"
    output_path = PROCESSADOS_DIR / output_name

    raw_text = subprocess.run(
        ["pdftotext", "-layout", str(pdf_path), "-"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout.replace("\f", "\n\n")

    cleaned_lines = [line.rstrip() for line in raw_text.splitlines()]
    cleaned_text = "\n".join(cleaned_lines).strip()
    if not cleaned_text:
        cleaned_text = "EXTRACAO_VAZIA_OU_BLOQUEADA_PELO_ARQUIVO_ORIGINAL"

    override = MANUAL_REVIEW_OVERRIDES.get(source_row["source_id"], {})
    review_status = override.get("status", "texto_extraido_layout_pendente_revisao_manual_tabelas")
    review_note = override.get("review_note")

    body = "\n".join(
        [
            "---",
            f"source_id: {source_row['source_id']}",
            f"titulo: {source_row['título']}",
            f"organizacao: {source_row['organização']}",
            f"tipo: {source_row['tipo']}",
            f"nivel_de_confiabilidade: {source_row['nível_de_confiabilidade']}",
            f"arquivo_original: {source_row['arquivo_local']}",
            f"arquivo_processado_logico: beach_handball_ai/fontes/05_processado/documentos_markdown/{output_name}",
            "texto_extraido_por: pdftotext -layout",
            f"status: {review_status}",
            f"data_processamento: {TODAY}",
            "---",
            "",
            f"# {source_row['source_id']}",
            "",
            f"## Titulo de origem",
            "",
            source_row["título"],
            "",
            *(
                [
                    "## Nota de revisao manual",
                    "",
                    review_note,
                    "",
                ]
                if review_note
                else []
            ),
            "## Texto extraido",
            "",
            cleaned_text,
            "",
        ]
    )
    output_path.write_text(body, encoding="utf-8")
    return output_name, f"beach_handball_ai/fontes/05_processado/documentos_markdown/{output_name}"


def required_artifact_rows() -> list[dict[str, str]]:
    return [
        {
            "source_id": "FONTES_OFICIAIS_CSV",
            "título": "Registro de Fontes Oficiais (CSV)",
            "organização": "ScoutPraia/CEPRAEA",
            "tipo": "Artefato de controle",
            "data": TODAY,
            "versão": TODAY,
            "link": "",
            "uso_permitido": "uso interno operacional",
            "tema": "catalogo/fontes",
            "nível_de_confiabilidade": "derivado",
            "observação": "Registro canônico em CSV do corpus combinado docs/sources + beach_handball_ai/fontes. O checksum desta linha e ciclico por definicao e deve ser auditado pelo manifest.",
            "arquivo_local": "beach_handball_ai/fontes/00_registro/fontes_oficiais.csv",
            "status": "registro_canonico_atualizado",
            "checksum_sha256": "",
            "data_coleta": TODAY,
            "source_code_repo": "",
            "origem_catalogo": "beach_handball_ai/fontes",
            "papel_documento": "artefato_de_controle",
            "fonte_primaria_relacionada": "",
            "uso_mvp": "sim",
            "uso_rag_fase2": "apoio_de_governanca",
        },
        {
            "source_id": "FONTES_OFICIAIS_XLSX",
            "título": "Registro de Fontes Oficiais (XLSX)",
            "organização": "ScoutPraia/CEPRAEA",
            "tipo": "Artefato de controle",
            "data": TODAY,
            "versão": TODAY,
            "link": "",
            "uso_permitido": "uso interno operacional",
            "tema": "catalogo/fontes",
            "nível_de_confiabilidade": "derivado",
            "observação": "Espelho operacional em planilha do registro canônico de fontes. O checksum desta linha e ciclico por definicao e deve ser auditado pelo manifest.",
            "arquivo_local": "beach_handball_ai/fontes/00_registro/fontes_oficiais.xlsx",
            "status": "registro_canonico_atualizado",
            "checksum_sha256": "",
            "data_coleta": TODAY,
            "source_code_repo": "",
            "origem_catalogo": "beach_handball_ai/fontes",
            "papel_documento": "artefato_de_controle",
            "fonte_primaria_relacionada": "FONTES_OFICIAIS_CSV",
            "uso_mvp": "sim",
            "uso_rag_fase2": "apoio_de_governanca",
        },
        {
            "source_id": "MANIFEST_CHECKSUMS_JSON",
            "título": "Manifest Checksums do Corpus",
            "organização": "ScoutPraia/CEPRAEA",
            "tipo": "Artefato de controle",
            "data": TODAY,
            "versão": TODAY,
            "link": "",
            "uso_permitido": "uso interno operacional",
            "tema": "processamento/manifesto",
            "nível_de_confiabilidade": "derivado",
            "observação": "Manifest SHA-256 do catalogo canonico e do corpus operacional. O checksum desta linha no registro e ciclico por dependencia mutua com CSV/XLSX.",
            "arquivo_local": "beach_handball_ai/fontes/05_processado/manifest_checksums.json",
            "status": "manifest_atualizado",
            "checksum_sha256": "",
            "data_coleta": TODAY,
            "source_code_repo": "",
            "origem_catalogo": "beach_handball_ai/fontes",
            "papel_documento": "artefato_de_controle",
            "fonte_primaria_relacionada": "",
            "uso_mvp": "sim",
            "uso_rag_fase2": "apoio_de_governanca",
        },
        {
            "source_id": "METADADOS_TEMATICOS_ETAPA_3_JSONL",
            "título": "Metadados Tematicos da Etapa 3",
            "organização": "ScoutPraia/CEPRAEA",
            "tipo": "Artefato de processamento",
            "data": TODAY,
            "versão": TODAY,
            "link": "",
            "uso_permitido": "uso interno operacional",
            "tema": "processamento/metadata",
            "nível_de_confiabilidade": "derivado",
            "observação": "Inventário pre-chunking com metadados temáticos para separação por tema.",
            "arquivo_local": "beach_handball_ai/fontes/05_processado/chunks_jsonl/METADADOS_TEMATICOS_ETAPA_3.jsonl",
            "status": "metadata_pre_chunking_atualizada",
            "checksum_sha256": "",
            "data_coleta": TODAY,
            "source_code_repo": "",
            "origem_catalogo": "beach_handball_ai/fontes",
            "papel_documento": "artefato_processado_consolidado",
            "fonte_primaria_relacionada": "",
            "uso_mvp": "nao",
            "uso_rag_fase2": "bloqueado_ate_fase2",
        },
        {
            "source_id": "CHUNKS_ETAPA_4_CORPUS_JSONL",
            "título": "Chunks Estruturados da Etapa 4",
            "organização": "ScoutPraia/CEPRAEA",
            "tipo": "Artefato de processamento",
            "data": TODAY,
            "versão": TODAY,
            "link": "",
            "uso_permitido": "uso interno operacional",
            "tema": "chunking/corpus_final",
            "nível_de_confiabilidade": "derivado",
            "observação": "Corpus estruturado de chunks da Etapa 4 com metadados por unidade, separacao de precedencia e historicos deprecated.",
            "arquivo_local": "beach_handball_ai/fontes/05_processado/chunks_jsonl/CHUNKS_ETAPA_4_CORPUS.jsonl",
            "status": "chunks_etapa_4_gerados",
            "checksum_sha256": "",
            "data_coleta": TODAY,
            "source_code_repo": "",
            "origem_catalogo": "beach_handball_ai/fontes",
            "papel_documento": "artefato_processado_consolidado",
            "fonte_primaria_relacionada": "",
            "uso_mvp": "nao",
            "uso_rag_fase2": "bloqueado_ate_fase2_global",
        },
        {
            "source_id": "CHUNKS_ETAPA_4_AUDITORIA_MD",
            "título": "Auditoria da Etapa 4",
            "organização": "ScoutPraia/CEPRAEA",
            "tipo": "Artefato de auditoria",
            "data": TODAY,
            "versão": TODAY,
            "link": "",
            "uso_permitido": "uso interno operacional",
            "tema": "chunking/auditoria",
            "nível_de_confiabilidade": "derivado",
            "observação": "Resumo de auditoria do corpus estruturado da Etapa 4 com contagens por status e tema.",
            "arquivo_local": "beach_handball_ai/fontes/05_processado/chunks_jsonl/CHUNKS_ETAPA_4_AUDITORIA.md",
            "status": "auditoria_chunking_etapa_4_atualizada",
            "checksum_sha256": "",
            "data_coleta": TODAY,
            "source_code_repo": "",
            "origem_catalogo": "beach_handball_ai/fontes",
            "papel_documento": "artefato_processado_consolidado",
            "fonte_primaria_relacionada": "CHUNKS_ETAPA_4_CORPUS_JSONL",
            "uso_mvp": "nao",
            "uso_rag_fase2": "bloqueado_ate_fase2_global",
        },
        {
            "source_id": "CHUNKS_FINAIS_RAG_IHF_RULES_BH_2026_PT_MD",
            "título": "CHUNKS_FINAIS_RAG_IHF_RULES_BH_2026_PT",
            "organização": "ScoutPraia/CEPRAEA",
            "tipo": "Artefato processado markdown",
            "data": TODAY,
            "versão": TODAY,
            "link": "",
            "uso_permitido": "uso interno operacional",
            "tema": "chunking/revisao_final",
            "nível_de_confiabilidade": "derivado",
            "observação": "Artefato legado de curadoria IHF convertido para o JSONL estruturado da Etapa 4; manter apenas como base historica auditavel.",
            "arquivo_local": "beach_handball_ai/fontes/05_processado/chunks_jsonl/CHUNKS_FINAIS_RAG_IHF_RULES_BH_2026_PT.md",
            "status": "historico_convertido_para_jsonl_etapa_4",
            "checksum_sha256": "",
            "data_coleta": TODAY,
            "source_code_repo": "",
            "origem_catalogo": "beach_handball_ai/fontes",
            "papel_documento": "artefato_processado_consolidado",
            "fonte_primaria_relacionada": "IHF_RULES_BH_2026_PT_TRANSLATION",
            "uso_mvp": "nao",
            "uso_rag_fase2": "bloqueado_ate_fase2_global",
        },
    ]


def build_auxiliary_processado_row(output_rel: str) -> dict[str, str]:
    path = Path(output_rel)
    stem = path.stem.upper()
    source_id = "AUX_PROCESSADO_" + "".join(char if char.isalnum() else "_" for char in stem).strip("_")
    tema = "processamento/auxiliar"
    tipo = "Artefato processado auxiliar"
    status = "historico_auxiliar_processo"
    observacao = "Artefato auxiliar/historico mantido para auditoria do fluxo de conversão e revisão."

    if "AUDITORIA_CORPUS" in stem:
        tema = "processamento/auditoria"
        tipo = "Artefato de auditoria do corpus"
        status = "auditoria_auxiliar_historica"
        observacao = "Artefato histórico de auditoria do corpus mantido para rastreabilidade; não é fonte primária nem artefato final de RAG."
    elif path.name == "audit_md_corpus.py":
        tema = "processamento/script_auxiliar"
        tipo = "Script auxiliar de auditoria"
        status = "script_auxiliar_historico"
        observacao = "Script auxiliar histórico mantido para rastreabilidade do fluxo de auditoria do corpus."
    elif "_rascunhos" in output_rel:
        tema = "processamento/rascunho"
        tipo = "Rascunho processado"
        status = "rascunho_historico_bloqueado"
        observacao = "Rascunho histórico de chunking/conversão mantido apenas para auditoria do processo; não liberar para RAG principal."

    return {
        "source_id": source_id,
        "título": path.stem,
        "organização": "ScoutPraia/CEPRAEA",
        "tipo": tipo,
        "data": TODAY,
        "versão": TODAY,
        "link": "",
        "uso_permitido": "uso interno operacional",
        "tema": tema,
        "nível_de_confiabilidade": "derivado",
        "observação": observacao,
        "arquivo_local": output_rel,
        "status": status,
        "checksum_sha256": "",
        "data_coleta": TODAY,
        "source_code_repo": "",
        "origem_catalogo": "beach_handball_ai/fontes",
        "papel_documento": "artefato_processado_auxiliar",
        "fonte_primaria_relacionada": "",
        "uso_mvp": "nao",
        "uso_rag_fase2": "bloqueado_ate_fase2",
    }


def sync_xlsx(rows: list[dict[str, str]]) -> None:
    wb = load_workbook(REGISTRY_XLSX)

    ws = wb["fontes_oficiais"]
    ws.delete_rows(1, ws.max_row)
    ws.append(FIELDNAMES)
    for row in rows:
        ws.append([row[field] for field in FIELDNAMES])

    ws = wb["crosswalk_repo"]
    ws.delete_rows(1, ws.max_row)
    crosswalk_header = [
        "source_code_repo",
        "source_id",
        "arquivo_local",
        "status",
        "papel_documento",
        "origem_catalogo",
    ]
    ws.append(crosswalk_header)
    for row in rows:
        if row["source_code_repo"]:
            ws.append([row[h] for h in crosswalk_header])

    ws = wb["registro_status"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["status", "quantidade"])
    for status, count in sorted(Counter(row["status"] for row in rows).items()):
        ws.append([status, count])

    ws = wb["documentos_processados"]
    ws.delete_rows(1, ws.max_row)
    docs_header = [
        "source_id",
        "titulo_processado",
        "url",
        "arquivo_processado_logico",
        "status",
        "checksum",
        "observacao",
        "data",
    ]
    ws.append(docs_header)
    for row in rows:
        if "05_processado/documentos_markdown" in row["arquivo_local"]:
            ws.append(
                [
                    row["fonte_primaria_relacionada"] or row["source_id"],
                    row["título"],
                    "",
                    row["arquivo_local"],
                    row["status"],
                    row["checksum_sha256"],
                    row["observação"],
                    row["data"] or row["versão"],
                ]
            )

    ws = wb["auditoria_organizacao"]
    ws.delete_rows(1, ws.max_row)
    header, audit_rows = build_auditoria_organizacao_rows(rows)
    ws.append(header)
    for row in audit_rows:
        ws.append(row)

    ws = wb["proxima_acao"]
    ws.delete_rows(1, ws.max_row)
    header, next_rows = build_proxima_acao_rows()
    ws.append(header)
    for row in next_rows:
        ws.append(row)

    ws = wb["revisao_cruzada"]
    ws.delete_rows(1, ws.max_row)
    header, review_rows = build_revisao_cruzada_rows()
    ws.append(header)
    for row in review_rows:
        ws.append(row)

    ws = wb["matriz_chunks_final"]
    ws.delete_rows(1, ws.max_row)
    header, chunk_rows = build_matriz_chunks_rows()
    ws.append(header)
    for row in chunk_rows:
        ws.append(row)

    wb.save(REGISTRY_XLSX)


def sync_manifest(rows: list[dict[str, str]]) -> None:
    files = []
    for root in CATALOG_ROOTS:
        for path in sorted(root.rglob("*")):
            if path.is_file():
                files.append(
                    {
                        "path": str(path.relative_to(ROOT)),
                        "size_bytes": path.stat().st_size,
                        "sha256": MANIFEST_SELF_SHA if path == MANIFEST_JSON else sha256(path),
                    }
                )

    manifest = {
        "generated_at": TODAY,
        "root": str(ROOT),
        "scope": ["docs/sources", "beach_handball_ai/fontes"],
        "notes": [
            "Registro unificado inclui catálogo canônico do repositório e corpus operacional de trabalho.",
            "RAG segue bloqueado pela política atual do ScoutPraia; este manifest não libera fase 2.",
            "PDFs catalogados foram convertidos para Markdown com pdftotext -layout em 05_processado/documentos_markdown quando aplicável.",
        ],
        "status_summary": dict(sorted(Counter(row["status"] for row in rows).items())),
        "files": files,
    }
    MANIFEST_JSON.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    rows = read_registry()
    rows_by_source = {row["source_id"]: row for row in rows}
    rows_by_file = {row["arquivo_local"]: row for row in rows}

    pdf_rows = [row for row in rows if row["arquivo_local"].lower().endswith(".pdf")]
    for source_row in pdf_rows:
        _, output_rel = extract_pdf_to_markdown(source_row)
        processed_source_id = f"MD_PROCESSADO_{source_row['source_id']}"
        rows_by_source[processed_source_id] = base_processed_row(source_row, output_rel)
        rows_by_file[output_rel] = rows_by_source[processed_source_id]

    sync_stage4_chunks(rows)

    for row in required_artifact_rows():
        rows_by_source[row["source_id"]] = row
        rows_by_file[row["arquivo_local"]] = row

    required_files = []
    for root in CATALOG_ROOTS:
        for path in sorted(root.rglob("*")):
            if path.is_file():
                required_files.append(str(path.relative_to(ROOT)))

    for path in required_files:
        if path not in rows_by_file and path.startswith("beach_handball_ai/fontes/05_processado/"):
            row = build_auxiliary_processado_row(path)
            rows_by_source[row["source_id"]] = row
            rows_by_file[path] = row

    missing_files = [path for path in required_files if path not in rows_by_file]
    if missing_files:
        raise SystemExit(f"Arquivos sem registro após refresh: {missing_files}")

    normalized_rows = [
        normalize_row(row)
        for row in rows_by_source.values()
        if (ROOT / row["arquivo_local"]).exists()
    ]
    write_registry(normalized_rows)
    sync_xlsx(normalized_rows)
    normalized_rows = [normalize_row(row) for row in read_registry()]
    write_registry(normalized_rows)
    sync_manifest(normalized_rows)
    normalized_rows = [normalize_row(row) for row in read_registry()]
    write_registry(normalized_rows)
    sync_xlsx(normalized_rows)


if __name__ == "__main__":
    main()
