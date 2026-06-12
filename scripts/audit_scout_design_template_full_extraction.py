"""Audit full chunk extraction of SCOUT_DESIGN_TEMPLATE.

This script validates whether a JSON/JSONL chunk export really represents the
spreadsheet in a way that agents/RAG can use safely.

It does not call Google Drive. Export the Google Doc/JSON view locally first or
produce a local JSON/JSONL file, then run this script.

Examples:
    PYTHONPATH=. python3 scripts/audit_scout_design_template_full_extraction.py \
        --chunks docs/SCOUT_DESIGN_TEMPLATE_FULL_EXTRACTION.json \
        --xlsx docs/SCOUT_DESIGN_TEMPLATE.xlsx

    PYTHONPATH=. python3 scripts/audit_scout_design_template_full_extraction.py \
        --chunks docs/SCOUT_DESIGN_TEMPLATE_AGENT_VIEW_CHUNKS.json
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - environment guard
    load_workbook = None  # type: ignore[assignment]


REQUIRED_CHUNK_FIELDS = frozenset(
    {
        "chunk_id",
        "title",
        "theme",
        "priority",
        "agent_use",
        "content",
    }
)

REQUIRED_FULL_EXTRACTION_FIELDS = REQUIRED_CHUNK_FIELDS | frozenset(
    {
        "sheet_name",
        "chunk_type",
        "row_range",
    }
)

CRITICAL_SHEETS = frozenset(
    {
        "MODULE_INDEX",
        "SHEET_MAP",
        "EVENTOS",
        "CROSS_MODULE_BOUNDARIES",
        "AI_USE_POLICY",
        "FIELD_DICTIONARY_GLOBAL",
        "RESULT_DOMAIN_GLOBAL",
        "VALIDATION_MATRIX",
        "LEGACY_MIGRATION_RULES",
        "SOURCE_REGISTER",
        "ARCHITECTURE_README",
        "EVENTOS_LEGADOS_FUTUROS",
        "EVENT_REQUIRED_FIELDS",
        "EVENT_OPTIONAL_FIELDS",
        "EVENT_FORBIDDEN_FIELDS",
        "EVENT_BLOCKING_RULES",
        "SCORER_ROLES",
        "PONTUACAO_FINALIZACAO",
        "DECISION_PRECEDENCE",
        "QUALITY_REQUIREMENTS",
        "ACCEPTANCE_EVIDENCE",
    }
)

REQUIRED_RULE_PATTERNS = {
    "specialist_is_not_event_code": r"specialist\s*(não|nao|not).*event_code|event_code\s*=\s*specialist_shot.*(proib|forbid|block)",
    "specialist_is_not_position_code": r"specialist\s*(não|nao|not).*position_code|position_code\s*=\s*specialist.*(proib|forbid|block)",
    "specialist_goal_two_points": r"simple_shot.*goal.*specialist.*2|specialist.*goal.*2",
    "shootout_save_not_goalkeeper_save": r"shootout.*save.*(não|nao|not).*goalkeeper_save|goalkeeper_save.*(não|nao|not).*shootout",
    "goalkeeper_save_requires_linked_finalization": r"goalkeeper_save.*linked_finalization_id|linked_finalization_id.*goalkeeper_save",
    "transition_does_not_score": r"transition_sequence.*(não|nao|not).*pontos|transição.*não calcula pontos|transition.*does not.*points",
    "transition_goal_requires_finalization_terminal": r"transition_goal.*finaliza|transition_goal.*terminal.*finalization|transition_goal.*Finalização",
    "g5_blocks_mvp_complete": r"G5.*pendente|MVP completo.*(não|nao|not).*declarado|MVP.*não pode.*completo",
    "rag_does_not_classify_lance": r"RAG.*(não|nao|not).*decide.*lance|RAG.*não decide lance|RAG.*não.*automaticamente",
}

ALLOWED_PRIORITIES = frozenset({"critical", "high", "medium", "low"})


@dataclass
class Finding:
    severity: str
    code: str
    message: str
    location: str = ""


@dataclass
class AuditReport:
    findings: list[Finding] = field(default_factory=list)

    @property
    def errors(self) -> list[Finding]:
        return [finding for finding in self.findings if finding.severity == "error"]

    @property
    def warnings(self) -> list[Finding]:
        return [finding for finding in self.findings if finding.severity == "warning"]

    @property
    def ok(self) -> bool:
        return not self.errors

    def add(self, severity: str, code: str, message: str, location: str = "") -> None:
        self.findings.append(Finding(severity=severity, code=code, message=message, location=location))


def load_chunks(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    if not path.exists():
        raise FileNotFoundError(path)

    text = path.read_text(encoding="utf-8").strip()
    if not text:
        return {}, []

    if path.suffix.lower() == ".jsonl":
        chunks = [json.loads(line) for line in text.splitlines() if line.strip()]
        return {"format": "jsonl"}, chunks

    payload = json.loads(text)
    if isinstance(payload, list):
        return {"format": "json_array"}, payload
    if isinstance(payload, dict):
        chunks = payload.get("chunks")
        if isinstance(chunks, list):
            return payload, chunks
        if {"chunk_id", "content"}.issubset(payload):
            return {"format": "single_chunk"}, [payload]
    raise ValueError("Arquivo precisa ser JSON com chave chunks, array JSON ou JSONL.")


def sheet_names_from_xlsx(path: Path) -> list[str]:
    if load_workbook is None:
        raise RuntimeError("openpyxl não está instalado; não é possível ler o XLSX.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def chunk_sheet_names(chunk: dict[str, Any]) -> set[str]:
    names: set[str] = set()
    sheet_name = chunk.get("sheet_name")
    if isinstance(sheet_name, str) and sheet_name.strip():
        names.add(sheet_name.strip())
    source_sheets = chunk.get("source_sheets") or chunk.get("source_sheet")
    if isinstance(source_sheets, str) and source_sheets.strip():
        names.add(source_sheets.strip())
    if isinstance(source_sheets, list):
        names.update(str(item).strip() for item in source_sheets if str(item).strip())
    return names


def serialize_for_search(payload: Any) -> str:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True)


def validate_chunk_structure(
    chunks: list[dict[str, Any]],
    report: AuditReport,
    *,
    full_extraction: bool,
) -> None:
    seen_ids: set[str] = set()
    required = REQUIRED_FULL_EXTRACTION_FIELDS if full_extraction else REQUIRED_CHUNK_FIELDS

    for index, chunk in enumerate(chunks):
        location = f"chunk[{index}]"
        if not isinstance(chunk, dict):
            report.add("error", "invalid_chunk_type", "Chunk precisa ser objeto JSON.", location)
            continue

        chunk_id = chunk.get("chunk_id")
        if not isinstance(chunk_id, str) or not chunk_id.strip():
            report.add("error", "missing_chunk_id", "Chunk sem chunk_id válido.", location)
        elif chunk_id in seen_ids:
            report.add("error", "duplicate_chunk_id", f"chunk_id duplicado: {chunk_id}", location)
        else:
            seen_ids.add(chunk_id)
            location = chunk_id

        missing_fields = sorted(field for field in required if field not in chunk or chunk[field] in (None, ""))
        if missing_fields:
            report.add(
                "error",
                "missing_required_chunk_fields",
                f"Campos obrigatórios ausentes: {', '.join(missing_fields)}",
                location,
            )

        priority = chunk.get("priority")
        if isinstance(priority, str) and priority not in ALLOWED_PRIORITIES:
            report.add(
                "warning",
                "unknown_priority",
                f"Prioridade fora do domínio recomendado: {priority}",
                location,
            )

        content = chunk.get("content")
        if content in (None, "", {}, []):
            report.add("error", "empty_chunk_content", "Chunk sem conteúdo útil.", location)


def validate_sheet_coverage(
    chunks: list[dict[str, Any]],
    expected_sheets: Iterable[str],
    report: AuditReport,
    *,
    require_all_sheets: bool,
) -> None:
    expected = set(expected_sheets)
    covered: set[str] = set()
    for chunk in chunks:
        covered.update(chunk_sheet_names(chunk))

    if require_all_sheets:
        missing = sorted(expected - covered)
        extra = sorted(covered - expected)
        if missing:
            report.add("error", "missing_sheet_coverage", f"Abas sem chunk: {', '.join(missing)}")
        if extra:
            report.add("warning", "unknown_sheet_referenced", f"Chunks referenciam abas não presentes no XLSX: {', '.join(extra)}")
    else:
        critical_missing = sorted((CRITICAL_SHEETS & expected) - covered)
        if critical_missing:
            report.add(
                "warning",
                "critical_sheet_not_referenced",
                f"Abas críticas não referenciadas por chunks sintéticos: {', '.join(critical_missing)}",
            )


def validate_critical_sheet_presence(chunks: list[dict[str, Any]], report: AuditReport) -> None:
    covered: set[str] = set()
    for chunk in chunks:
        covered.update(chunk_sheet_names(chunk))
    if not covered:
        report.add(
            "warning",
            "no_sheet_level_metadata",
            "Nenhum chunk informa sheet_name/source_sheets; isso é aceitável só para visão sintética, não para extração completa.",
        )
        return
    missing_critical = sorted(CRITICAL_SHEETS - covered)
    if missing_critical:
        report.add(
            "warning",
            "critical_sheets_missing_from_chunks",
            f"Abas críticas sem chunk direto: {', '.join(missing_critical)}",
        )


def validate_required_rules(metadata: dict[str, Any], chunks: list[dict[str, Any]], report: AuditReport) -> None:
    search_text = serialize_for_search({"metadata": metadata, "chunks": chunks})
    for rule_id, pattern in REQUIRED_RULE_PATTERNS.items():
        if re.search(pattern, search_text, flags=re.IGNORECASE | re.DOTALL):
            continue
        report.add("error", "missing_required_rule", f"Regra crítica ausente dos chunks: {rule_id}")


def audit_extraction(
    *,
    chunks_path: Path,
    xlsx_path: Path | None = None,
    require_all_sheets: bool = False,
    full_extraction: bool = False,
) -> AuditReport:
    report = AuditReport()
    try:
        metadata, chunks = load_chunks(chunks_path)
    except Exception as exc:  # noqa: BLE001 - audit CLI should report user-facing errors
        report.add("error", "cannot_load_chunks", f"Não foi possível carregar chunks: {exc}")
        return report

    if not chunks:
        report.add("error", "no_chunks", "Nenhum chunk encontrado.")
        return report

    validate_chunk_structure(chunks, report, full_extraction=full_extraction)
    validate_required_rules(metadata, chunks, report)
    validate_critical_sheet_presence(chunks, report)

    expected_sheets: list[str] = []
    if xlsx_path is not None:
        try:
            expected_sheets = sheet_names_from_xlsx(xlsx_path)
        except Exception as exc:  # noqa: BLE001
            report.add("error", "cannot_read_xlsx", f"Não foi possível ler XLSX: {exc}")
            expected_sheets = []

    if expected_sheets:
        validate_sheet_coverage(
            chunks,
            expected_sheets,
            report,
            require_all_sheets=require_all_sheets or full_extraction,
        )

    declared_sheet_count = metadata.get("expected_sheet_count") if isinstance(metadata, dict) else None
    if isinstance(declared_sheet_count, int) and expected_sheets and declared_sheet_count != len(expected_sheets):
        report.add(
            "error",
            "declared_sheet_count_mismatch",
            f"expected_sheet_count={declared_sheet_count}, XLSX tem {len(expected_sheets)} abas.",
        )

    return report


def format_report(report: AuditReport) -> str:
    lines = [
        "== SCOUT_DESIGN_TEMPLATE full extraction audit ==",
        f"status={'ok' if report.ok else 'failed'}",
        f"errors={len(report.errors)}",
        f"warnings={len(report.warnings)}",
    ]
    for label, findings in (("errors", report.errors), ("warnings", report.warnings)):
        if not findings:
            continue
        lines.append(f"\n-- {label} --")
        for finding in findings:
            location = f" {finding.location}" if finding.location else ""
            lines.append(f"- [{finding.code}]{location} {finding.message}")
    return "\n".join(lines)


def parse_args(argv: Iterable[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Audita chunks extraídos do SCOUT_DESIGN_TEMPLATE")
    parser.add_argument("--chunks", required=True, type=Path, help="Arquivo JSON/JSONL com chunks")
    parser.add_argument("--xlsx", type=Path, default=None, help="SCOUT_DESIGN_TEMPLATE.xlsx para validar cobertura de abas")
    parser.add_argument(
        "--require-all-sheets",
        action="store_true",
        help="Exige que todas as abas do XLSX estejam cobertas por chunks",
    )
    parser.add_argument(
        "--full-extraction",
        action="store_true",
        help="Exige campos de extração completa: sheet_name, chunk_type e row_range",
    )
    return parser.parse_args(list(argv))


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    report = audit_extraction(
        chunks_path=args.chunks,
        xlsx_path=args.xlsx,
        require_all_sheets=args.require_all_sheets,
        full_extraction=args.full_extraction,
    )
    print(format_report(report))
    return 0 if report.ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
