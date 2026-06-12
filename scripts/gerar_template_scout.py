#!/usr/bin/env python3
"""Gera ou regenera docs/SCOUT_DESIGN_TEMPLATE.xlsx com todas as abas de governança.

Escopo:
    - Cria o workbook completo a partir do zero via openpyxl
    - Não lê o XLSX existente — sobrescreve completamente

Modo: MUTANTE — sobrescreve docs/SCOUT_DESIGN_TEMPLATE.xlsx
Gate/trigger: quando a estrutura de abas do template precisa ser recriada do zero
Artefatos produzidos: docs/SCOUT_DESIGN_TEMPLATE.xlsx

Examples:
    python3 scripts/gerar_template_scout.py
"""

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# ── cores ──────────────────────────────────────────────────────────────────────
AZUL_HEADER   = "1F4E79"  # cabeçalho
AZUL_SECAO    = "2E75B6"  # título de seção
CINZA_HEADER  = "D6DCE4"  # sub-cabeçalho
VERDE_NOVO    = "E2EFDA"  # linha nova / a preencher
AMARELO_REV   = "FFF2CC"  # linha existente que precisa revisar
BRANCO        = "FFFFFF"
CINZA_CLARO   = "F2F2F2"

# ── helpers ────────────────────────────────────────────────────────────────────

def cabecalho(ws, colunas: list[tuple[str, int]]):
    """Escreve linha de cabeçalho com cores e ajusta largura."""
    row = []
    for titulo, largura in colunas:
        row.append(titulo)
    ws.append(row)
    for col_idx, (_, largura) in enumerate(colunas, start=1):
        cel = ws.cell(row=ws.max_row, column=col_idx)
        cel.font = Font(bold=True, color="FFFFFF", size=10)
        cel.fill = PatternFill("solid", fgColor=AZUL_HEADER)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = largura
    ws.row_dimensions[ws.max_row].height = 28


def linha(ws, valores: list, cor: str = BRANCO, negrito: bool = False):
    ws.append(valores)
    for col_idx, _ in enumerate(valores, start=1):
        cel = ws.cell(row=ws.max_row, column=col_idx)
        cel.fill = PatternFill("solid", fgColor=cor)
        cel.font = Font(size=9, bold=negrito)
        cel.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 40


def linha_vazia(ws, ncols: int, cor: str = VERDE_NOVO):
    ws.append([""] * ncols)
    for col_idx in range(1, ncols + 1):
        cel = ws.cell(row=ws.max_row, column=col_idx)
        cel.fill = PatternFill("solid", fgColor=cor)
    ws.row_dimensions[ws.max_row].height = 36


def titulo_secao(ws, texto: str, ncols: int):
    ws.append([texto] + [""] * (ncols - 1))
    cel = ws.cell(row=ws.max_row, column=1)
    cel.font = Font(bold=True, color="FFFFFF", size=10)
    cel.fill = PatternFill("solid", fgColor=AZUL_SECAO)
    ws.merge_cells(
        start_row=ws.max_row, start_column=1,
        end_row=ws.max_row, end_column=ncols
    )
    ws.row_dimensions[ws.max_row].height = 22


def nota(ws, texto: str, ncols: int):
    ws.append([texto] + [""] * (ncols - 1))
    cel = ws.cell(row=ws.max_row, column=1)
    cel.font = Font(italic=True, size=9, color="595959")
    ws.merge_cells(
        start_row=ws.max_row, start_column=1,
        end_row=ws.max_row, end_column=ncols
    )
    ws.row_dimensions[ws.max_row].height = 18


# ── aba INSTRUÇÕES ─────────────────────────────────────────────────────────────

def aba_instrucoes(wb):
    ws = wb.create_sheet("INSTRUÇÕES")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 120

    linhas = [
        ("TEMPLATE DE DESIGN DO SCOUT — ScoutPraia", True,  AZUL_HEADER, "FFFFFF"),
        ("", False, BRANCO, "000000"),
        ("OBJETIVO", True, AZUL_SECAO, "FFFFFF"),
        ("Este arquivo é o documento de design do instrumento de scout.",
         False, BRANCO, "000000"),
        ("Você preenche aqui, com seu conhecimento técnico, TODOS os campos antes de qualquer implementação.",
         False, BRANCO, "000000"),
        ("Quando o design estiver fechado, a IA implementa tudo de uma vez.",
         False, BRANCO, "000000"),
        ("", False, BRANCO, "000000"),
        ("COMO USAR", True, AZUL_SECAO, "FFFFFF"),
        ("1. Aba EVENTOS: defina cada ação que você quer marcar no vídeo.",
         False, CINZA_CLARO, "000000"),
        ("2. Aba POSIÇÕES: posição tática da atleta no momento da ação.",
         False, BRANCO, "000000"),
        ("3. Aba ZONAS_QUADRA: de onde na quadra a ação acontece.",
         False, CINZA_CLARO, "000000"),
        ("4. Aba ZONAS_GOL: para onde a bola vai dentro do gol (só finalizações).",
         False, BRANCO, "000000"),
        ("5. Aba SISTEMAS: sistema tático em que a jogada ocorreu.",
         False, CINZA_CLARO, "000000"),
        ("6. Aba RESULTADOS_POSSE: como uma posse pode terminar.",
         False, BRANCO, "000000"),
        ("", False, BRANCO, "000000"),
        ("LEGENDA DE CORES", True, AZUL_SECAO, "FFFFFF"),
        ("Linha BRANCA → já existe no sistema atual",
         False, BRANCO, "000000"),
        ("Linha AMARELA → existe, mas precisa revisão ou complemento",
         False, AMARELO_REV, "000000"),
        ("Linha VERDE → campo novo, ainda não implementado — você preenche",
         False, VERDE_NOVO, "000000"),
        ("", False, BRANCO, "000000"),
        ("CAMPOS OBRIGATÓRIOS vs OPCIONAIS", True, AZUL_SECAO, "FFFFFF"),
        ("Sim   → obrigatório para salvar o evento",
         False, CINZA_CLARO, "000000"),
        ("Não   → não se aplica a esse evento",
         False, BRANCO, "000000"),
        ("Opc   → recomendado mas não bloqueia o salvamento",
         False, CINZA_CLARO, "000000"),
        ("", False, BRANCO, "000000"),
        ("DICA OPERACIONAL", True, AZUL_SECAO, "FFFFFF"),
        ("Priorize os campos que você realmente vai olhar no relatório.",
         False, BRANCO, "000000"),
        ("Campo que você não usa no relatório não precisa existir na tela.",
         False, BRANCO, "000000"),
        ("Se um campo te faz hesitar mais de 2s no vídeo, repense a definição.",
         False, BRANCO, "000000"),
    ]

    for texto, neg, fundo, fonte_cor in linhas:
        ws.append([texto])
        cel = ws.cell(row=ws.max_row, column=1)
        cel.font = Font(bold=neg, size=10 if neg else 9, color=fonte_cor)
        cel.fill = PatternFill("solid", fgColor=fundo)
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 20 if not neg else 24


# ── aba EVENTOS ────────────────────────────────────────────────────────────────

def aba_eventos(wb):
    ws = wb.create_sheet("EVENTOS")
    ws.freeze_panes = "A2"

    cols = [
        ("codigo",              18),
        ("nome_ui",             24),
        ("categoria",           18),
        ("definicao",           40),
        ("marcar_quando",       32),
        ("nao_marcar_quando",   32),
        ("regra_decisao",       32),
        ("pontos_possiveis",    14),
        ("req_atleta",          10),
        ("req_zona_quadra",     12),
        ("req_zona_gol",        12),
        ("req_posicao",         10),
        ("req_sistema",         10),
        ("status",              10),
        ("observacao",          36),
    ]
    cabecalho(ws, cols)
    N = len(cols)

    # ── OFENSIVOS ──────────────────────────────────────────────────────────────
    titulo_secao(ws, "OFENSIVOS", N)
    nota(ws, "Ações de ataque da SUA equipe", N)

    eventos_ofensivos = [
        ("shot_attempt",      "Tentativa de Finalização", "Ofensivo",
         "arremesso com intenção clara de gol",
         "há arremesso ou ação clara de finalização",
         "passe, finta sem arremesso",
         "se a intenção principal é finalizar, marcar",
         "0", "Sim", "Opc", "Sim", "Opc", "existente", ""),

        ("goal_scored",       "Gol Marcado",              "Ofensivo",
         "finalização que resulta em gol válido de 1 ponto",
         "placar sobe 1 para a equipe",
         "gol anulado, gol de 2 pontos",
         "confirmar pontos_value=1",
         "1", "Sim", "Sim", "Opc", "Opc", "existente", ""),

        ("shot_missed",       "Finalização Fora/Trave",   "Ofensivo",
         "finalização sem gol e sem defesa da goleira",
         "bola vai fora ou na trave sem intervenção da goleira",
         "defesa clara da goleira",
         "se a goleira alterar a trajetória preferir save",
         "0", "Sim", "Sim", "Opc", "Opc", "existente", ""),

        ("technical_error",   "Erro Técnico",             "Ofensivo",
         "perda de posse sem arremesso por erro técnico ou decisão",
         "passe errado, recepção falha, condução/violação",
         "arremesso defendido, roubo claro",
         "sem arremesso e sem ação defensiva clara → erro técnico",
         "0", "Sim", "Não", "Opc", "Opc", "existente", ""),

        ("turnover",          "Perda de Posse",           "Ofensivo",
         "qualquer perda de posse antes de nova posse própria",
         "adversária passa a controlar",
         "fim de set, gol marcado",
         "usar como categoria ampla; subtipo explica causa",
         "0", "Opc", "Não", "Opc", "Opc", "existente", ""),

        ("assist",            "Assistência",              "Ofensivo",
         "passe imediatamente relacionado ao gol",
         "passe cria finalização convertida",
         "passe anterior sem relação direta",
         "marcar atleta secundária como assistente",
         "0", "Sim", "Não", "Opc", "Opc", "existente", ""),

        ("spin_shot",         "Arremesso Spin",           "Ofensivo",
         "finalização com giro característico",
         "atleta executa giro antes da finalização",
         "finta com giro sem arremesso",
         "se houver dúvida, registrar como two_point_attempt e revisar",
         "0 ou 2", "Sim", "Sim", "Opc", "Opc", "existente", ""),

        ("inflight_attempt",  "Tentativa Inflight",       "Ofensivo",
         "atleta recebe/controla no ar e finaliza antes de tocar o solo",
         "finalização ocorre no ar após passe/recepção",
         "passe alto sem finalização, finalização após contato solo",
         "contar apenas com tentativa de finalização",
         "0 ou 2", "Sim", "Sim", "Opc", "Opc", "existente", ""),

        ("inflight_goal",     "Gol Inflight",             "Ofensivo",
         "inflight_attempt convertido em gol válido",
         "finalização aérea resulta em gol",
         "gol comum, gol anulado",
         "points_value deve refletir regra aplicável (geralmente 2)",
         "2", "Sim", "Sim", "Opc", "Opc", "existente", ""),

        ("two_point_attempt", "Tentativa de 2 Pontos (genérico)", "Ofensivo",
         "tentativa de ação que pode valer 2 pontos — categoria genérica",
         "quando não for spin, inflight, especialista ou shoot-out claramente identificados",
         "quando o tipo específico já está identificado",
         "prefira o tipo específico quando possível",
         "0 ou 2", "Sim", "Sim", "Opc", "Opc", "revisar",
         "Considere se ainda faz sentido manter esse genérico"),

        ("two_point_goal",    "Gol de 2 Pontos (genérico)", "Ofensivo",
         "gol válido de 2 pontos — categoria genérica",
         "quando o tipo específico não está claro",
         "gol comum, gol anulado",
         "points_value=2 obrigatório",
         "2", "Sim", "Sim", "Opc", "Opc", "revisar",
         "Considere se ainda faz sentido manter esse genérico"),
    ]

    for ev in eventos_ofensivos:
        cor = AMARELO_REV if ev[-2] == "revisar" else BRANCO
        linha(ws, list(ev), cor)

    # ── ESPECIALISTA (NOVOS) ───────────────────────────────────────────────────
    titulo_secao(ws, "ESPECIALISTA — preencha conforme sua regra operacional", N)
    nota(ws, "Gol da especialista = 2 pontos por qualquer arremesso (Regra 9:6). Ainda não está implementado.", N)

    especialista = [
        ("specialist_attempt", "Tentativa da Especialista", "Ofensivo",
         "", "", "", "", "0 ou 2", "Sim", "Sim", "Sim", "Opc", "novo", ""),
        ("specialist_goal",    "Gol da Especialista",       "Ofensivo",
         "", "", "", "", "2",     "Sim", "Sim", "Sim", "Opc", "novo", ""),
    ]
    for ev in especialista:
        linha_vazia_dados = list(ev)
        ws.append(linha_vazia_dados)
        for col_idx, _ in enumerate(linha_vazia_dados, start=1):
            cel = ws.cell(row=ws.max_row, column=col_idx)
            cel.fill = PatternFill("solid", fgColor=VERDE_NOVO)
            cel.font = Font(size=9)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 40

    nota(ws, "▶ Adicione abaixo outros eventos ofensivos que ainda não existem", N)
    for _ in range(3):
        linha_vazia(ws, N, VERDE_NOVO)

    # ── DEFENSIVOS ─────────────────────────────────────────────────────────────
    titulo_secao(ws, "DEFENSIVOS", N)
    nota(ws, "Ações defensivas da SUA equipe", N)

    eventos_def = [
        ("defensive_stop",      "Parada Defensiva",       "Defensivo",
         "posse adversária encerrada sem gol por ação defensiva relevante",
         "defesa força arremesso ruim, perda ou interrupção favorável",
         "erro adversário sem pressão clara",
         "marcar quando a defesa altera claramente a qualidade da posse",
         "0", "Opc", "Não", "Opc", "Opc", "existente", ""),

        ("steal",               "Roubo de Bola",          "Defensivo",
         "recuperação direta de posse por ação defensiva",
         "atleta intercepta ou toma a bola",
         "bola perdida sem controle defensivo imediato",
         "precisa haver ganho de posse claro",
         "0", "Sim", "Não", "Opc", "Opc", "existente", ""),

        ("block",               "Bloqueio",               "Defensivo",
         "bloqueio de arremesso pela defesa de linha",
         "defensor altera ou impede trajetória do arremesso",
         "defesa da goleira",
         "se a goleira é responsável, usar save",
         "0", "Sim", "Não", "Opc", "Opc", "existente", ""),

        ("forced_error",        "Erro Forçado",           "Defensivo",
         "erro adversário causado por pressão defensiva clara",
         "pressão força passe ruim, violação ou perda",
         "erro sem pressão identificável",
         "se não houver evidência clara de pressão, não marcar",
         "0", "Opc", "Não", "Opc", "Opc", "existente", ""),

        ("goal_conceded",       "Gol Sofrido",            "Defensivo",
         "gol sofrido pela equipe",
         "adversária marca gol válido",
         "gol anulado",
         "registrar valor em pontos concedidos",
         "variável", "Opc", "Não", "Opc", "Opc", "existente", ""),

        ("defensive_breakdown", "Falha Defensiva",        "Defensivo",
         "falha coletiva ou individual que gera finalização clara adversária",
         "desorganização, troca perdida, cobertura ausente",
         "gol sofrido sem falha identificável",
         "manter em testing — interpretativo",
         "0", "Opc", "Não", "Opc", "Opc", "existente",
         "Evento interpretativo — calibrar com vídeo"),
    ]
    for ev in eventos_def:
        cor = AMARELO_REV if ev[-2] == "revisar" else BRANCO
        linha(ws, list(ev), cor)

    nota(ws, "▶ Adicione abaixo outros eventos defensivos que ainda não existem", N)
    for _ in range(3):
        linha_vazia(ws, N, VERDE_NOVO)

    # ── GOLEIRA ────────────────────────────────────────────────────────────────
    titulo_secao(ws, "GOLEIRA", N)

    eventos_gol = [
        ("save",                    "Defesa da Goleira",        "Goleira",
         "goleira altera ou impede gol em arremesso adversário",
         "goleira toca e desvia ou segura a bola",
         "bloqueio de linha, bola na trave sem toque",
         "marcar quando há intervenção clara da goleira",
         "0", "Sim", "Não", "Não", "Não", "existente", ""),

        ("save_shootout",           "Defesa no Shoot-out",      "Goleira",
         "defesa da goleira especificamente em shoot-out",
         "só usar em situação de shoot-out",
         "defesa em jogo normal",
         "",
         "0", "Sim", "Não", "Não", "Não", "existente", ""),

        ("goalkeeper_distribution", "Reposição da Goleira",     "Goleira",
         "reposição/passe da goleira que inicia ataque",
         "marcar quando gera posse organizada ou vantagem clara",
         "reposição simples sem vantagem",
         "",
         "0", "Sim", "Não", "Opc", "Não", "existente",
         "Importante para transição — goleira como primeiro passe do ataque"),
    ]
    for ev in eventos_gol:
        linha(ws, list(ev), BRANCO)

    nota(ws, "▶ Adicione abaixo outros eventos de goleira que ainda não existem", N)
    for _ in range(2):
        linha_vazia(ws, N, VERDE_NOVO)

    # ── TRANSIÇÃO ──────────────────────────────────────────────────────────────
    titulo_secao(ws, "TRANSIÇÃO", N)

    eventos_trans = [
        ("fast_break_for",            "Contra-ataque Favorável",      "Transição",
         "transição ofensiva rápida da equipe após recuperação",
         "posse rápida gera finalização ou vantagem clara",
         "ataque posicional após posse",
         "",
         "0", "Opc", "Não", "Opc", "Opc", "existente", ""),

        ("fast_break_against",        "Contra-ataque Adversário",     "Transição",
         "transição ofensiva rápida da adversária",
         "adversária finaliza ou cria vantagem clara em transição",
         "ataque posicional adversário",
         "",
         "0", "Opc", "Não", "Opc", "Opc", "existente", ""),

        ("transition_recovery_good",  "Recomposição Defensiva Boa",   "Transição",
         "equipe recupera posição defensiva eficientemente",
         "equipe impede vantagem clara da adversária em transição",
         "recomposição com gol sofrido",
         "",
         "0", "Opc", "Não", "Opc", "Opc", "existente", ""),

        ("transition_recovery_bad",   "Recomposição Defensiva Ruim",  "Transição",
         "falha de recomposição permite finalização adversária",
         "adversária obtém finalização clara por atraso defensivo",
         "falha que não gera finalização clara",
         "",
         "0", "Opc", "Não", "Opc", "Opc", "existente", ""),
    ]
    for ev in eventos_trans:
        linha(ws, list(ev), BRANCO)

    nota(ws, "▶ Adicione abaixo outros eventos de transição que ainda não existem", N)
    for _ in range(2):
        linha_vazia(ws, N, VERDE_NOVO)

    # ── SHOOT-OUT ──────────────────────────────────────────────────────────────
    titulo_secao(ws, "SHOOT-OUT", N)

    eventos_so = [
        ("shootout_attempt", "Tentativa Shoot-out", "Shoot-out",
         "tentativa em shoot-out — marcar todas, convertidas ou não",
         "toda tentativa de shoot-out",
         "",
         "",
         "0 ou 2", "Sim", "Sim", "Não", "Não", "existente", ""),

        ("shootout_goal",    "Gol no Shoot-out",    "Shoot-out",
         "shoot-out convertido em gol",
         "bola entra no gol",
         "defesa, trave, fora",
         "points_value conforme regra",
         "2", "Sim", "Sim", "Não", "Não", "existente", ""),

        ("shootout_miss",    "Shoot-out Errado",    "Shoot-out",
         "shoot-out não convertido sem defesa da goleira",
         "bola vai fora ou na trave",
         "defesa da goleira → usar save_shootout",
         "separar defesa da goleira com save_shootout",
         "0", "Sim", "Sim", "Não", "Não", "existente", ""),
    ]
    for ev in eventos_so:
        linha(ws, list(ev), BRANCO)

    # ── SITUAÇÕES ESPECIAIS ────────────────────────────────────────────────────
    titulo_secao(ws, "SITUAÇÕES ESPECIAIS", N)

    eventos_esp = [
        ("timeout",  "Tempo Técnico", "Situação Especial",
         "pedido de tempo registrado no vídeo",
         "marcar timestamp para contexto do relatório",
         "", "",
         "0", "Não", "Não", "Não", "Não", "existente", ""),

        ("set_end",  "Fim do Set",    "Situação Especial",
         "fim do set",
         "marcar placar e timestamp final",
         "", "",
         "0", "Não", "Não", "Não", "Não", "existente", ""),
    ]
    for ev in eventos_esp:
        linha(ws, list(ev), BRANCO)

    nota(ws, "▶ Adicione abaixo outros eventos de situação especial", N)
    for _ in range(2):
        linha_vazia(ws, N, VERDE_NOVO)


# ── aba POSIÇÕES ───────────────────────────────────────────────────────────────

def aba_posicoes(wb):
    ws = wb.create_sheet("POSIÇÕES")
    ws.freeze_panes = "A2"

    cols = [
        ("codigo",      16),
        ("nome_ui",     24),
        ("contexto",    20),
        ("descricao",   50),
        ("status",      12),
        ("observacao",  36),
    ]
    cabecalho(ws, cols)
    N = len(cols)

    nota(ws, "Posição tática da atleta NO MOMENTO da ação. Preencha com as posições que você usa no seu sistema.", N)

    titulo_secao(ws, "POSIÇÕES EXISTENTES NO SISTEMA ATUAL", N)
    nota(ws, "O sistema atual não tem posições. Todas abaixo são NOVAS — defina conforme seu modelo de jogo.", N)

    posicoes = [
        ("goalkeeper",   "Goleira",              "Defensivo / Ofensivo como especialista",
         "jogadora na função de goleira", "novo", ""),
        ("specialist",   "Especialista",         "Ofensivo",
         "jogadora com colete da especialista — qualquer gol vale 2pts (Regra 9:6)", "novo",
         "Fundamental para calcular eficiência da especialista separadamente"),
        ("line_player_1","Jogadora de Linha 1",  "Ofensivo / Defensivo",
         "", "novo", "Renomeie conforme sua terminologia"),
        ("line_player_2","Jogadora de Linha 2",  "Ofensivo / Defensivo",
         "", "novo", "Renomeie conforme sua terminologia"),
    ]

    for p in posicoes:
        ws.append(list(p))
        for col_idx in range(1, N + 1):
            cel = ws.cell(row=ws.max_row, column=col_idx)
            cel.fill = PatternFill("solid", fgColor=VERDE_NOVO)
            cel.font = Font(size=9)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 40

    nota(ws, "▶ Adicione abaixo as posições que faltam no seu modelo de jogo", N)
    for _ in range(4):
        linha_vazia(ws, N, VERDE_NOVO)


# ── aba ZONAS_QUADRA ───────────────────────────────────────────────────────────

def aba_zonas_quadra(wb):
    ws = wb.create_sheet("ZONAS_QUADRA")
    ws.freeze_panes = "A2"

    cols = [
        ("codigo",      18),
        ("nome_ui",     22),
        ("descricao",   44),
        ("contexto",    20),
        ("status",      12),
        ("observacao",  36),
    ]
    cabecalho(ws, cols)
    N = len(cols)

    nota(ws, "Zona da quadra onde a ação acontece. Revise se a divisão atual faz sentido para o seu modelo de análise.", N)

    titulo_secao(ws, "ZONAS EXISTENTES", N)

    zonas = [
        ("left_wing",      "Ponta Esquerda",     "extremidade esquerda da quadra", "Ofensivo / Defensivo", "existente", ""),
        ("left_half",      "Meia Esquerda",       "setor esquerdo de meio",          "Ofensivo / Defensivo", "existente", ""),
        ("center",         "Centro",              "setor central",                   "Ofensivo / Defensivo", "existente", ""),
        ("right_half",     "Meia Direita",        "setor direito de meio",           "Ofensivo / Defensivo", "existente", ""),
        ("right_wing",     "Ponta Direita",       "extremidade direita da quadra",   "Ofensivo / Defensivo", "existente", ""),
        ("6m_left",        "6m Esquerda",         "arremesso de 6m lado esquerdo",   "Ofensivo / Defensivo", "existente", ""),
        ("6m_center",      "6m Centro",           "arremesso de 6m central",         "Ofensivo / Defensivo", "existente", ""),
        ("6m_right",       "6m Direita",          "arremesso de 6m lado direito",    "Ofensivo / Defensivo", "existente", ""),
        ("shootout_lane",  "Corredor Shoot-out",  "linha de shoot-out",              "Shoot-out",            "existente", ""),
    ]
    for z in zonas:
        linha(ws, list(z), BRANCO)

    titulo_secao(ws, "ZONAS A REVISAR / ADICIONAR", N)
    nota(ws, "O modelo atual não distingue posição de ataque vs posição de arremesso. Revise se precisa de mais granularidade.", N)
    for _ in range(4):
        linha_vazia(ws, N, VERDE_NOVO)


# ── aba ZONAS_GOL ──────────────────────────────────────────────────────────────

def aba_zonas_gol(wb):
    ws = wb.create_sheet("ZONAS_GOL")
    ws.freeze_panes = "A2"

    cols = [
        ("codigo",      18),
        ("nome_ui",     26),
        ("linha_gol",   14),
        ("coluna_gol",  14),
        ("descricao",   36),
        ("status",      12),
        ("observacao",  36),
    ]
    cabecalho(ws, cols)
    N = len(cols)

    nota(ws, "NOVO — Onde a bola vai dentro do gol. Grade 3x3 (linha: Alto/Meio/Baixo × coluna: Esquerdo/Centro/Direito). Preencha nomes.", N)

    titulo_secao(ws, "GRADE 3×3 DO GOL — preencha nome_ui e observações", N)

    grade = [
        ("goal_top_left",    "Alto Esquerdo",    "Alto",  "Esquerdo", "", "novo", ""),
        ("goal_top_center",  "Alto Centro",      "Alto",  "Centro",   "", "novo", ""),
        ("goal_top_right",   "Alto Direito",     "Alto",  "Direito",  "", "novo", ""),
        ("goal_mid_left",    "Meio Esquerdo",    "Meio",  "Esquerdo", "", "novo", ""),
        ("goal_mid_center",  "Meio Centro",      "Meio",  "Centro",   "", "novo", ""),
        ("goal_mid_right",   "Meio Direito",     "Meio",  "Direito",  "", "novo", ""),
        ("goal_bot_left",    "Baixo Esquerdo",   "Baixo", "Esquerdo", "", "novo", ""),
        ("goal_bot_center",  "Baixo Centro",     "Baixo", "Centro",   "", "novo", ""),
        ("goal_bot_right",   "Baixo Direito",    "Baixo", "Direito",  "", "novo", ""),
    ]
    for g in grade:
        ws.append(list(g))
        for col_idx in range(1, N + 1):
            cel = ws.cell(row=ws.max_row, column=col_idx)
            cel.fill = PatternFill("solid", fgColor=VERDE_NOVO)
            cel.font = Font(size=9)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 36

    nota(ws, "▶ Você pode usar uma grade diferente (2x3, 3x4, etc.) — ajuste os códigos conforme quiser", N)
    nota(ws, "▶ Decide: quer capturar zona do gol APENAS em finalizações? Ou também em defesas da goleira?", N)


# ── aba SISTEMAS ───────────────────────────────────────────────────────────────

def aba_sistemas(wb):
    ws = wb.create_sheet("SISTEMAS")
    ws.freeze_panes = "A2"

    cols = [
        ("codigo",       20),
        ("nome_ui",      28),
        ("tipo",         16),
        ("descricao",    44),
        ("quando_usar",  36),
        ("status",       12),
        ("observacao",   36),
    ]
    cabecalho(ws, cols)
    N = len(cols)

    nota(ws, "NOVO — Sistema tático em que a jogada ocorreu. Preencha conforme seu modelo de jogo.", N)

    titulo_secao(ws, "OFENSIVOS — preencha com seus sistemas de ataque", N)
    sistemas_of = [
        ("positional_attack",    "Ataque Posicional",          "Ofensivo", "", "", "novo", ""),
        ("fast_break_attack",    "Contra-ataque",              "Ofensivo", "", "", "novo", ""),
        ("specialist_on_court",  "Especialista em Quadra",     "Ofensivo",
         "ataque com jogadora de linha substituindo a goleira como especialista",
         "", "novo", ""),
        ("empty_goal_attack",    "Ataque Gol Vazio",           "Ofensivo",
         "goleira fora — ataque com superioridade numérica",
         "", "novo", ""),
    ]
    for s in sistemas_of:
        ws.append(list(s))
        for col_idx in range(1, N + 1):
            cel = ws.cell(row=ws.max_row, column=col_idx)
            cel.fill = PatternFill("solid", fgColor=VERDE_NOVO)
            cel.font = Font(size=9)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 40

    nota(ws, "▶ Adicione abaixo outros sistemas ofensivos", N)
    for _ in range(3):
        linha_vazia(ws, N, VERDE_NOVO)

    titulo_secao(ws, "DEFENSIVOS — preencha com seus sistemas de defesa", N)
    sistemas_def = [
        ("man_to_man",   "Marcação Individual", "Defensivo", "", "", "novo", ""),
        ("zone_defense", "Defesa em Zona",       "Defensivo", "", "", "novo", ""),
        ("high_press",   "Pressão Alta",         "Defensivo", "", "", "novo", ""),
    ]
    for s in sistemas_def:
        ws.append(list(s))
        for col_idx in range(1, N + 1):
            cel = ws.cell(row=ws.max_row, column=col_idx)
            cel.fill = PatternFill("solid", fgColor=VERDE_NOVO)
            cel.font = Font(size=9)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 40

    nota(ws, "▶ Adicione abaixo outros sistemas defensivos", N)
    for _ in range(3):
        linha_vazia(ws, N, VERDE_NOVO)


# ── aba RESULTADOS_POSSE ───────────────────────────────────────────────────────

def aba_resultados_posse(wb):
    ws = wb.create_sheet("RESULTADOS_POSSE")
    ws.freeze_panes = "A2"

    cols = [
        ("codigo",          20),
        ("nome_ui",         28),
        ("descricao",       44),
        ("pontos_equipe",   14),
        ("status",          12),
        ("observacao",      36),
    ]
    cabecalho(ws, cols)
    N = len(cols)

    nota(ws, "Como uma posse pode terminar. O campo 'resultado' da posse vira texto livre hoje — aqui você define os valores controlados.", N)

    titulo_secao(ws, "RESULTADOS ATUAIS (texto livre) → transformar em lista controlada", N)

    resultados = [
        ("goal",             "Gol",                  "posse termina em gol da equipe",                "1 ou 2", "novo", ""),
        ("technical_error",  "Erro Técnico",          "posse termina por erro técnico próprio",         "0", "novo", ""),
        ("turnover",         "Perda de Posse",        "adversária recupera sem erro técnico claro",     "0", "novo", ""),
        ("save",             "Defesa da Goleira",     "goleira adversária defende a finalização",       "0", "novo", ""),
        ("shot_wide",        "Finalização Fora",      "bola sai sem intervenção da goleira",            "0", "novo", ""),
        ("shot_blocked",     "Finalização Bloqueada", "bloqueio de linha impede a finalização",         "0", "novo", ""),
        ("forced_error",     "Erro Forçado",          "erro causado por pressão defensiva adversária",  "0", "novo", ""),
        ("shootout_won",     "Shoot-out Conquistado", "falta recebe shoot-out",                         "0→2", "novo", ""),
        ("goal_conceded",    "Gol Sofrido",           "adversária converte finalização",                "0", "novo", ""),
        ("penalty_conceded", "Shoot-out Sofrido",     "equipe comete falta, adversária ganha shoot-out","0→2", "novo", ""),
    ]
    for r in resultados:
        ws.append(list(r))
        for col_idx in range(1, N + 1):
            cel = ws.cell(row=ws.max_row, column=col_idx)
            cel.fill = PatternFill("solid", fgColor=VERDE_NOVO)
            cel.font = Font(size=9)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 36

    nota(ws, "▶ Adicione abaixo outros desfechos de posse que você usa", N)
    for _ in range(3):
        linha_vazia(ws, N, VERDE_NOVO)


# ── main ───────────────────────────────────────────────────────────────────────

def main():
    wb = Workbook()
    wb.remove(wb.active)  # remove aba padrão vazia

    aba_instrucoes(wb)
    aba_eventos(wb)
    aba_posicoes(wb)
    aba_zonas_quadra(wb)
    aba_zonas_gol(wb)
    aba_sistemas(wb)
    aba_resultados_posse(wb)

    caminho = "docs/SCOUT_DESIGN_TEMPLATE.xlsx"
    wb.save(caminho)
    print(f"Template gerado: {caminho}")


if __name__ == "__main__":
    main()
